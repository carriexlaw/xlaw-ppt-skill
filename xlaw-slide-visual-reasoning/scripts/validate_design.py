#!/usr/bin/env python3
"""
validate_design.py — 按 references/14-validation-spec.md 校验 deck

    python validate_design.py deck.pptx deck.manifest.yaml [--thresholds thresholds.yaml] [--json _qa/validate.json] [--qa _qa]
    python validate_design.py --list          # 列出全部 C-xx、级别、阈值来源

流程（§2）：schema → roles → geometry → deck → images → crosscheck → composite。
每项输出 {page, check, level, message, values}；任一 M 失败 → exit 1；W 只汇总。
所有数值阈值来自 thresholds.yaml（§8）。§7 里出现、但 §8 未列的一个值（IMAGE_LAYOUT_TOL）以代码常量给出并在 --list 中标明。
每个检查函数签名统一 (ctx, page) -> list[Finding]，deck 级检查 page=None。
层级模型（第五轮）：第一层级 = 本页论述对象，可以是一组（hero:* 0–4 个 / 一组 heading / 图表 / 表格 / 时间线），manifest 用 focus 声明；校验同组一致（C-01）与层级完整性（C-03 [W]）。
留白模型（尺度优先）：留白不是输入。校验尺度（S-01）、间距 g（S-02）、内容块位置（S-03）、容器贴内容（S-04）、无洞（S-05）、对齐线（S-06）。
"""
import argparse
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field, asdict
from typing import Optional, List, Dict, Any

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from roles import classify_slide, parse_role  # noqa: E402
from inkbox import text_ink_box, text_line_count, cjk_equiv, EMU_PER_PT  # noqa: E402

A_NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'

# §3.5 / §7 里出现但 §8 未列的值（--list 中标为「代码常量」）
IMAGE_LAYOUT_TOL = 0.05              # §7：配图方式反算时形状框比例的容差（相对页宽 / 页高）

ENUM = {
    'type': {'cover', 'agenda', 'section', 'content', 'closing', 'statement', 'quote'},   # quote = statement 的旧名
    'focus.form': {'numbers', 'headings', 'chart', 'table', 'timeline', 'image', 'statement'},
    'visual_side': {'left', 'right', 'full', 'none'},
    'density': {'light', 'medium', 'heavy'},
    'container': {'none', 'divider', 'fill', 'stroke'},
    'contrast': {'strong', 'restrained'},
    'title_pos': {'base', 'mask-edge', 'below-image', 'right-of-image', 'image-center'},
    'image.role': {'semantic', 'atmosphere', 'full-bleed', 'small'},
    'image.side': {'left', 'right', 'top', 'bottom', 'full'},
    'tone': {'argument', 'statement', 'vision'},
    'image.source': {'pexels', 'unsplash', 'pixabay', 'user'},
}
STOCK_SOURCES = {'pexels', 'unsplash', 'pixabay'}
CONTENT_REQUIRED = ['message', 'pattern', 'focus', 'density', 'container', 'columns', 'contrast',
                    'title_pos', 'subtitle', 'conclusion', 'series', 'g', 'source_chars']
IMAGE_REQUIRED = ['role', 'layout', 'side', 'source', 'id', 'photographer', 'url', 'query',
                  'candidates_viewed', 'reason']
DECK_REQUIRED = ['page', 'margins', 'title_anchor', 'pagenum_anchor', 'type_scale', 'title_size',
                 'min_size', 'fonts', 'fonts_latin', 'palette', 'accent_max_chars', 'shape_language', 'tone']
EXCLUDED_ROLES = {'bg', 'pagenum'}
EXCLUDED_IMAGE_QUALS = {'full-bleed'}          # 氛围图在 foreground 内（参与间距 / 内容块 / 对齐线），但不计元素数、不作主元素
NON_COUNTABLE = {'title', 'kicker', 'conclusion', 'source', 'arrow', 'deco'}
COLUMNS_RE = re.compile(r'\d+(:\d+)*')
# 紧贴对（§3.4b）：这些角色对贴在一起（间距 < 0.8g 或相交）时合并成一个元素再算间距、洞与对齐线
TIGHT_PAIRS = {frozenset(p) for p in (('hero', 'label'), ('hero', 'arrow'), ('icon', 'body'), ('icon', 'heading'), ('legend',),
                                      ('tag', 'card'), ('tag', 'body'), ('tag', 'heading'), ('tag', 'arrow'))}
# title_pos 与 image.layout / side 的合法组合（C-02）
LEGAL_TITLE_POS = {
    (1, None): {'base'}, (2, None): {'base'}, (7, None): {'base'}, (3, None): {'mask-edge'},
    (4, 'top'): {'below-image'}, (4, 'bottom'): {'base'},
    (5, 'left'): {'right-of-image'}, (5, 'right'): {'base'},
    (6, 'left'): {'right-of-image', 'image-center'}, (6, 'right'): {'base', 'image-center'},
}


# ---------------------------------------------------------------------------- 基础类型

@dataclass
class Finding:
    page: Optional[int]
    check: str
    level: str            # M / W
    message: str
    values: Dict[str, Any] = field(default_factory=dict)


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float

    @property
    def right(self): return self.x + self.w
    @property
    def bottom(self): return self.y + self.h
    @property
    def cx(self): return self.x + self.w / 2
    @property
    def cy(self): return self.y + self.h / 2
    @property
    def area(self): return max(self.w, 0) * max(self.h, 0)

    def contains(self, o, tol=0.0):
        return (self.x - tol <= o.x and self.y - tol <= o.y and
                o.right <= self.right + tol and o.bottom <= self.bottom + tol)

    def intersects(self, o, eps=0.0):
        """eps > 0 时相切也算相交（连通判定：各扩半径后恰好相切即为连通）"""
        return not (o.x > self.right + eps or o.right < self.x - eps or o.y > self.bottom + eps or o.bottom < self.y - eps) if eps else \
            not (o.x >= self.right or o.right <= self.x or o.y >= self.bottom or o.bottom <= self.y)

    def intersection(self, o):
        x1, y1 = max(self.x, o.x), max(self.y, o.y)
        x2, y2 = min(self.right, o.right), min(self.bottom, o.bottom)
        return Box(x1, y1, max(0, x2 - x1), max(0, y2 - y1))

    def pad(self, p):
        return Box(self.x - p, self.y - p, self.w + 2 * p, self.h + 2 * p)

    def r(self):
        return {'x': round(self.x, 1), 'y': round(self.y, 1), 'w': round(self.w, 1), 'h': round(self.h, 1)}


def union_box(boxes):
    x1 = min(b.x for b in boxes); y1 = min(b.y for b in boxes)
    x2 = max(b.right for b in boxes); y2 = max(b.bottom for b in boxes)
    return Box(x1, y1, x2 - x1, y2 - y1)


@dataclass
class Run:
    text: str
    size: Optional[float]
    fonts: List[str]          # rPr 里出现的 typeface；空 = 继承
    color: Optional[str]      # hex 大写 / 'scheme:xx' / None(继承)
    bold: bool
    para: int
    latin: Optional[str] = None     # a:latin typeface
    spc: float = 0.0                # 字间距 pt（rPr spc / 100）


@dataclass
class ShapeInfo:
    shape: Any
    name: str
    role: str
    qual: str
    box: Box
    ink: Box
    is_text: bool
    is_picture: bool
    is_table: bool
    is_chart: bool
    runs: List[Run]
    fill: Optional[str]       # hex / None(noFill) / 'gradient' / 'scheme:xx'
    fill_alpha: float         # 0–1 不透明度
    line: Optional[str]
    line_w: float
    n_rows: int

    @property
    def full_role(self):
        return f'{self.role}:{self.qual}' if self.qual else self.role

    @property
    def text(self):
        return ''.join(r.text for r in self.runs)

    @property
    def is_image_role(self):
        return self.role == 'image' or (self.role == 'hero' and self.qual == 'image')

    @property
    def image_qual(self):
        """图片角色限定词；hero:image 视为 semantic"""
        if self.role == 'image':
            return self.qual
        if self.role == 'hero' and self.qual == 'image':
            return 'semantic'
        return None


# ---------------------------------------------------------------------------- XML 读取工具

def _color_of(fill_el):
    """solidFill 子节点 → (hex|'scheme:x', alpha)"""
    srgb = fill_el.find(A_NS + 'srgbClr')
    if srgb is not None:
        alpha = srgb.find(A_NS + 'alpha')
        a = int(alpha.get('val')) / 100000.0 if alpha is not None else 1.0
        return srgb.get('val').upper(), a
    sch = fill_el.find(A_NS + 'schemeClr')
    if sch is not None:
        alpha = sch.find(A_NS + 'alpha')
        a = int(alpha.get('val')) / 100000.0 if alpha is not None else 1.0
        return 'scheme:' + sch.get('val'), a
    return None, 1.0


def _fill_of(sppr):
    """spPr → (fill, alpha)"""
    if sppr is None:
        return None, 1.0
    if sppr.find(A_NS + 'noFill') is not None:
        return None, 1.0
    sf = sppr.find(A_NS + 'solidFill')
    if sf is not None:
        return _color_of(sf)
    if sppr.find(A_NS + 'gradFill') is not None:
        return 'gradient', 1.0
    if sppr.find(A_NS + 'pattFill') is not None:
        return 'pattern', 1.0
    return None, 1.0


def _line_of(sppr):
    if sppr is None:
        return None, 0.0
    ln = sppr.find(A_NS + 'ln')
    if ln is None or ln.find(A_NS + 'noFill') is not None:
        return None, 0.0
    w = int(ln.get('w')) / EMU_PER_PT if ln.get('w') else 0.75
    sf = ln.find(A_NS + 'solidFill')
    if sf is not None:
        c, _ = _color_of(sf)
        return c, w
    return None, 0.0


def _runs_of_txbody(txbody):
    out = []
    if txbody is None:
        return out
    for pi, p in enumerate(txbody.findall(A_NS + 'p')):
        for child in p:
            if child.tag not in (A_NS + 'r', A_NS + 'fld'):
                continue
            t = child.find(A_NS + 't')
            text = (t.text or '') if t is not None else ''
            rpr = child.find(A_NS + 'rPr')
            size, fonts, color, bold, latin, spc = None, [], None, False, None, 0.0
            if rpr is not None:
                if rpr.get('sz'):
                    size = int(rpr.get('sz')) / 100.0
                bold = rpr.get('b') == '1'
                if rpr.get('spc'):
                    try:
                        spc = int(rpr.get('spc')) / 100.0
                    except ValueError:
                        spc = 0.0
                lt = rpr.find(A_NS + 'latin')
                if lt is not None:
                    latin = lt.get('typeface')
                for tag in ('latin', 'ea', 'cs'):
                    f = rpr.find(A_NS + tag)
                    if f is not None and f.get('typeface'):
                        fonts.append(f.get('typeface'))
                sf = rpr.find(A_NS + 'solidFill')
                if sf is not None:
                    color, _ = _color_of(sf)
            out.append(Run(text, size, fonts, color, bold, pi, latin, spc))
    return out


def shape_info(sh, th):
    role, qual = parse_role(sh.name)
    box = Box(sh.left / EMU_PER_PT, sh.top / EMU_PER_PT, sh.width / EMU_PER_PT, sh.height / EMU_PER_PT)
    has_tf = bool(getattr(sh, 'has_text_frame', False) and sh.has_text_frame)
    is_table = bool(getattr(sh, 'has_table', False) and sh.has_table)
    is_chart = bool(getattr(sh, 'has_chart', False) and sh.has_chart)
    is_picture = sh.shape_type is not None and int(sh.shape_type) == 13  # MSO_SHAPE_TYPE.PICTURE
    runs, n_rows = [], 0
    if has_tf:
        runs = _runs_of_txbody(sh.text_frame._txBody)
    # pptxgenjs 给纯形状也写一个空 txBody：没有文字的形状按形状框算墨迹框，不算文字形状
    is_text = has_tf and any(r.text.strip() for r in runs)
    if is_table:
        n_rows = len(sh.table.rows)
        for row in sh.table.rows:
            for cell in row.cells:
                runs.extend(_runs_of_txbody(cell.text_frame._txBody))
    sppr = sh._element.find('.//' + P_NS + 'spPr')
    if sppr is None:
        sppr = sh._element.find(P_NS + 'spPr')
    fill, alpha = _fill_of(sppr)
    line, lw = _line_of(sppr)
    # 带填充 / 描边的形状（tag、card、panel、色块里写字）的墨迹框 = 形状框；只有纯文本框才按文字估
    ink = Box(*text_ink_box(sh, th)) if (is_text and role not in ('tag', 'card', 'panel') and not fill and not line) else box
    return ShapeInfo(sh, sh.name or '', role, qual, box, ink, is_text, is_picture, is_table, is_chart,
                     runs, fill, alpha, line, lw, n_rows)


def _table_cell_fills(si):
    out = []
    if not si.is_table:
        return out
    for tc in si.shape._element.iter(A_NS + 'tc'):
        tcpr = tc.find(A_NS + 'tcPr')
        if tcpr is not None:
            sf = tcpr.find(A_NS + 'solidFill')
            if sf is not None:
                out.append(_color_of(sf)[0])
            ln_tags = [c for c in tcpr if c.tag.startswith(A_NS + 'ln')]
            for ln in ln_tags:
                sf = ln.find(A_NS + 'solidFill')
                if sf is not None:
                    out.append(_color_of(sf)[0])
    return out


# ---------------------------------------------------------------------------- 光栅化

# ---------------------------------------------------------------------------- 上下文

class PageCtx:
    def __init__(self, deck, idx, slide, mf):
        self.deck = deck
        self.idx = idx
        self.slide = slide
        self.mf = mf or {}
        self.type = self.mf.get('type', 'content')
        self.th = deck.th
        self.roles, self.role_errors = classify_slide(slide, self.type, self.th.get('hierarchy', {}).get('max_heroes', 4))
        self.shapes = [shape_info(sh, self.th) for sh in slide.shapes]
        self.by_role = defaultdict(list)
        for s in self.shapes:
            self.by_role[s.role].append(s)
        self._cache = {}
        self.results = {}       # 供 crosscheck / composite 读取的反算结果

    # ---- 集合
    @property
    def is_content(self):
        return self.type == 'content'

    def excluded(self, s):
        return (s.role in EXCLUDED_ROLES or s.role == 'mask' or (s.role == 'image' and s.qual in EXCLUDED_IMAGE_QUALS)
                or (s.role == 'panel' and s.qual == 'region'))          # 区域背景像 bg 一样排除

    @property
    def foreground(self):
        return [s for s in self.shapes if not self.excluded(s)]

    @property
    def countable(self):
        return [s for s in self.foreground if s.role not in NON_COUNTABLE and s.image_qual != 'atmosphere']

    @property
    def title(self):
        t = self.by_role.get('title', [])
        return t[0] if t else None

    @property
    def kicker(self):
        k = self.by_role.get('kicker', [])
        return k[0] if k else None

    @property
    def hero(self):
        h = self.by_role.get('hero', [])
        return h[0] if h else None

    @property
    def heroes(self):
        return self.by_role.get('hero', [])

    @property
    def regions(self):
        return [s for s in self.by_role.get('panel', []) if s.qual == 'region']

    @property
    def images(self):
        return [s for s in self.shapes if s.is_image_role]

    @property
    def masks(self):
        return self.by_role.get('mask', [])

    @property
    def cards(self):
        return self.by_role.get('card', [])

    # ---- 区域
    @property
    def margin_box(self):
        d = self.deck
        return Box(d.margins['l'], d.margins['t'], d.page_w - d.margins['l'] - d.margins['r'],
                   d.page_h - d.margins['t'] - d.margins['b'])

    @property
    def title_bottom(self):
        ys = [s.ink.bottom for s in (self.title, self.kicker) if s is not None]
        return max(ys) if ys else None

    @property
    def body_box(self):
        """B = 版心 − 标题区 − 3g（title 下沿到 B.top 恰为 3g，见 C-09）"""
        m = self.margin_box
        tb = self.title_bottom
        if tb is None:
            return m
        top = tb + self.title_gap()
        return Box(m.x, top, m.w, max(m.bottom - top, 0))

    # ---- 容器与层
    @property
    def containers(self):
        return self.cards + [s for s in self.by_role.get('panel', []) if s.qual != 'region']

    def host_of(self, s):
        """包含 s 的最小容器；无则 None（顶层）"""
        tol = self.th['scale']['align_tol']
        hosts = [c for c in self.containers if c is not s and c.box.contains(s.box, tol)]
        return min(hosts, key=lambda c: c.box.area) if hosts else None

    def children(self, c):
        return [s for s in self.foreground if self.host_of(s) is c]

    def top_level(self, shapes=None):
        return [s for s in (shapes if shapes is not None else self.foreground) if self.host_of(s) is None]

    def is_edge_image(self, s):
        """配图方式 5 / 6：满页高、贴左或右页边的图片。它是版面的一条边，不进内容块的竖向范围与对齐线"""
        d = self.deck
        return bool(s.is_image_role and s.box.y <= 1 and s.box.bottom >= d.page_h - 1 and s.box.w < d.page_w - 2
                    and (s.box.x <= 1 or s.box.right >= d.page_w - 1))

    def content_shapes(self):
        """内容块成员：foreground 减 title / kicker（pagenum 已在 excluded），且不在标题区之上"""
        tb = self.title_bottom
        out = []
        for s in self.foreground:
            if s.role in ('title', 'kicker') or self.is_edge_image(s):
                continue
            if tb is not None and s.box.bottom <= tb:
                continue
            out.append(s)
        return out

    def content_box(self):
        cs = self.content_shapes()
        return union_box([s.ink for s in cs]) if cs else None

    def columns(self):
        """顶层内容形状（形状框）按 x 投影重叠 ≥ 30%（以较窄者计）传递聚类成列；返回 [(union_box, [shapes])]"""
        cs = self.top_level(self.content_shapes())
        ov = self.th['scale']['overlap']
        parent = list(range(len(cs)))

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i
        for i in range(len(cs)):
            for j in range(i + 1, len(cs)):
                a, b = cs[i].box, cs[j].box
                o = min(a.right, b.right) - max(a.x, b.x)
                if o >= ov * min(a.w, b.w):
                    parent[find(i)] = find(j)
        groups = defaultdict(list)
        for i, sh in enumerate(cs):
            groups[find(i)].append(sh)
        cols = [(union_box([sh.box for sh in g]), g) for g in groups.values()]
        return sorted(cols, key=lambda c: c[0].x)

    @property
    def bg_color(self):
        bg = self.by_role.get('bg', [])
        if bg and bg[0].fill and not bg[0].fill.startswith('scheme'):
            return bg[0].fill
        return 'FFFFFF'

    # ---- 文本统计
    def body_size_counter(self):
        """正文文本（role = body）的 run 字号计数（按字符数加权）；hero / heading / label / legend / tag / conclusion 不参与"""
        c = Counter()
        for s in self.countable:
            if s.role != 'body':
                continue
            for r in s.runs:
                if r.size and r.text.strip():
                    c[r.size] += len(r.text)
        return c

    def body_font_size(self):
        """正文字号：本页正文文本（role = body）的众数；本页无正文文本时取 deck 级正文字号"""
        c = self.body_size_counter()
        if c:
            return c.most_common(1)[0][0]
        return self.deck.body_font_size()

    def density_chars(self):
        n = 0
        for s in self.foreground:
            if s.role in ('title', 'pagenum', 'source'):
                continue
            n += cjk_equiv(s.text)
        return n

    def element_count(self):
        # 以元素为单位（§3.4b：紧贴对合并成一个元素）；元素里只要有一个 countable 形状就计 1
        #   - icon：作为要点符号或某个条目配图的 icon 与该条目合计 1 个，不单独计数（紧贴对已合并；
        #     单独成元素的 icon 只要与某个非 icon 的计数元素同行或同列（投影重叠 ≥ 30%）就算它的配图）
        #   - 并列组（S-06 的几何判定：同层、同角色、同上沿 / 中线、等间距，n ≥ 3）按 1 + n / 4 计，与 table 一致
        cnt = {id(s) for s in self.countable}
        ov = self.th['scale']['overlap']
        us = [u for u in self.units() if any(id(s) in cnt for s in u['shapes'])]

        def icon_only(u):
            return all(s.role == 'icon' for s in u['shapes'])

        def attached(u):
            for v in us:
                if v is u or icon_only(v) or v['layer'] is not u['layer']:
                    continue
                a, b = u['box'], v['box']
                ox = min(a.right, b.right) - max(a.x, b.x)
                oy = min(a.bottom, b.bottom) - max(a.y, b.y)
                if ox >= ov * min(a.w, b.w) or oy >= ov * min(a.h, b.h):
                    return True
            return False
        us = [u for u in us if not (icon_only(u) and attached(u))]
        n, grouped = 0.0, set()
        tol = self.th['scale']['align_tol']
        for layer in [None] + self.containers:
            rows = []
            _align_edges([u for u in us if u['layer'] is layer], tol, rows_out=rows, keep_units=True)
            for row in rows:
                if len(row) >= 3:
                    n += 1 + len(row) / 4
                    grouped.update(id(u) for u in row)
        for u in us:
            if id(u) in grouped:
                continue
            ms = [s for s in u['shapes'] if id(s) in cnt]
            n += max((1 + s.n_rows / 4) if s.is_table else 1 for s in ms)
        return n

    # ---- 紧贴对 → 元素（§3.4b）
    def units(self):
        """foreground 形状按紧贴对并查集合并；返回 [Unit]，Unit = {shapes, box, ink, layer}"""
        if 'units' in self._cache:
            return self._cache['units']
        fg = self.foreground
        tight = self.th['scale'].get('tight_max_mult', 0.8) * self.g()
        ov = self.th['scale']['overlap']
        layer = {id(s): self.host_of(s) for s in fg}
        parent = list(range(len(fg)))

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i
        for i in range(len(fg)):
            for j in range(i + 1, len(fg)):
                a, b = fg[i], fg[j]
                if frozenset((a.role, b.role)) not in TIGHT_PAIRS:
                    continue
                if 'card' in (a.role, b.role):
                    if not a.box.intersects(b.box):          # 胶囊 / 圆形标题压在框的角上
                        continue
                elif layer[id(a)] is not layer[id(b)]:
                    continue
                else:
                    ai, bi = a.ink, b.ink
                    o2 = min(ai.bottom, bi.bottom) - max(ai.y, bi.y)
                    hgap = max(a.box.x, b.box.x) - min(a.box.right, b.box.right)       # 横向用形状框（生成端按内容收紧了框）
                    o1 = min(ai.right, bi.right) - max(ai.x, bi.x)
                    vgap = max(ai.y, bi.y) - min(ai.bottom, bi.bottom)
                    side = o2 >= ov * min(ai.h, bi.h) and hgap <= tight
                    stack = frozenset((a.role, b.role)) == frozenset(('tag', 'arrow')) and o1 >= ov * min(ai.w, bi.w) and vgap <= tight
                    if not (side or stack):
                        continue
                parent[find(i)] = find(j)
        groups = defaultdict(list)
        for i, sh in enumerate(fg):
            groups[find(i)].append(sh)
        out = []
        for g_ in groups.values():
            host = next((layer[id(x)] for x in g_ if x.role != 'tag'), layer[id(g_[0])])
            out.append({'shapes': g_, 'box': union_box([x.box for x in g_]), 'ink': union_box([x.ink for x in g_]), 'layer': host,
                        'sig': tuple(sorted(x.role for x in g_))})
        self._cache['units'] = out
        return out

    def row_mates(self, u, us, region_of, strict=False):
        """与 u 同一行的元素：同层、同一区域背景范围、横向不重叠，且顶沿 / 垂直中心对齐（±2pt）或竖向墨迹重叠 ≥ 30%（以较矮者计）。
        一格里上下叠了两个元素（如「30% → 60%」+ 说明）时，它们都和旁边居中的 icon / 小标题同行"""
        tol, ov = self.th['scale']['align_tol'], self.th['scale']['overlap']
        ru = region_of(u)
        out = []
        for v in us:
            if v is u:
                out.append(v); continue
            if v['layer'] is not u['layer'] or region_of(v) is not ru:
                continue
            if not (v['box'].x >= u['box'].right - 0.5 or v['box'].right <= u['box'].x + 0.5):
                continue
            o = min(u['ink'].bottom, v['ink'].bottom) - max(u['ink'].y, v['ink'].y)
            if abs(v['box'].y - u['box'].y) <= tol or abs(v['box'].cy - u['box'].cy) <= tol or (not strict and o >= ov * min(u['ink'].h, v['ink'].h) > 0):
                out.append(v)
        return out

    # ---- 相邻间距与 g（§3.4），缓存
    def gaps(self):
        if 'gaps' in self._cache:
            return self._cache['gaps']
        ov = self.th['scale']['overlap']
        tol = self.th['scale']['align_tol']
        us = [u for u in self.units() if not any(self.is_edge_image(x) for x in u['shapes'])]
        # 行：同层、顶沿或垂直中心对齐（±2pt）且横向不重叠的元素。竖向间距在行与行之间量（行高由最高的那格决定）
        col_boxes = [c for c, _ in self.columns()]
        regions = self.regions

        def region_of(u):
            return next((r for r in regions if r.box.contains(u['box'], tol)), None)
        row_of = {id(u): self.row_mates(u, us, region_of) for u in us}
        strict_of = {id(u): self.row_mates(u, us, region_of, strict=True) for u in us}      # 只按对齐判的行：用来定对方那一行的横向范围
        out = []      # (value, kind, a, b)：a / b 取元素里的第一个形状（报错用）
        for ua in us:
            best_v, best_h = None, None
            a_row = row_of[id(ua)]
            for ub in us:
                if ua is ub or ua['layer'] is not ub['layer']:
                    continue
                if ua['box'].contains(ub['box']) or ub['box'].contains(ua['box']):
                    continue
                ai, bi = ua['ink'], ub['ink']
                ab, bb = ua['box'], ub['box']
                o = min(ab.right, bb.right) - max(ab.x, bb.x)          # 上下相邻按形状框判（文本框宽即列宽），距离按墨迹框量
                if o >= ov * min(ab.w, bb.w) and bi.y >= ai.bottom - 0.5 and not any(ub is v for v in a_row):
                    g = bi.y - ai.bottom
                    b_row = row_of[id(ub)]
                    bx0, bx1 = min(v['box'].x for v in strict_of[id(ub)]), max(v['box'].right for v in strict_of[id(ub)])
                    ax0, ax1 = min(v['box'].x for v in strict_of[id(ua)]), max(v['box'].right for v in strict_of[id(ua)])
                    # 只取与对方那一行横向有交叠的同行元素（并排的两个独立格子各管各的）
                    row_bottom = max(v['ink'].bottom for v in a_row if v is ua or (v['box'].x < bx1 and v['box'].right > bx0))
                    row_top = min(v['ink'].y for v in b_row if v is ub or (v['box'].x < ax1 and v['box'].right > ax0))
                    rb, ra = region_of(ub), region_of(ua)
                    if rb is not None and ra is not rb:                 # 从区域背景外走进去：量到区域内容的上沿
                        row_top = min(v['ink'].y for v in us if region_of(v) is rb)
                    alt = row_top - row_bottom if row_top - row_bottom >= -0.5 else g      # 行到行
                    if best_v is None or g < best_v[0]:
                        best_v = (g, 'v', ua['shapes'][0], ub['shapes'][0], None, None, alt)
                o2 = min(ai.bottom, bi.bottom) - max(ai.y, bi.y)
                if o2 >= ov * min(ai.h, bi.h) and ub['box'].x >= ua['box'].right - 0.5:
                    if all(x.role in ('tag', 'icon') for x in ua['shapes']) or all(x.role in ('tag', 'icon') for x in ub['shapes']):
                        continue                        # 小容器按文字收宽、icon 按自身尺寸，都比所在的格子窄；横向间距由它所在的列 / 同组的文字判定
                    blocked = any(uc is not ua and uc is not ub and uc['layer'] is ua['layer'] and uc['box'].x >= ua['box'].right - 0.5 and uc['box'].right <= ub['box'].x + 0.5
                                  and min(uc['ink'].bottom, ai.bottom) - max(uc['ink'].y, ai.y) > 0 for uc in us)
                    if not blocked and ua['layer'] is None:
                        blocked = any(c.x >= ua['box'].right - 0.5 and c.right <= ub['box'].x + 0.5 for c in col_boxes)     # 中间隔着一整列（本行在那一列是空位）
                    if blocked:
                        continue                        # 中间隔着另一列：不是相邻
                    g = ub['box'].x - ua['box'].right   # 横向用形状框：文本框宽即列宽，墨迹右沿随对齐与行长变化
                    rb, ra = region_of(ub), region_of(ua)
                    if rb is not None and ra is not rb:          # 横向走进区域背景：量到区域内最靠左的元素（区域里收窄居中的文字 / 图不改变分栏间距）
                        g = min(v['box'].x for v in us if region_of(v) is rb) - ua['box'].right
                    elif ra is not None and rb is not ra:
                        g = ub['box'].x - max(v['box'].right for v in us if region_of(v) is ra)
                    if best_h is None or g < best_h[0]:
                        best_h = (g, 'h', ua['shapes'][0], ub['shapes'][0], ua, ub)
            for cand in (best_v, best_h):
                if cand:
                    out.append(cand)
        self._cache['gaps'] = out
        return out

    def g(self):
        """g 由 layout.py 算出并写进页级 manifest；特殊页（无 g）取正文字号"""
        v = self.mf.get('g')
        if isinstance(v, (int, float)) and v > 0:
            return float(v)
        return float(self.body_font_size() or self.deck.th['scale']['bands']['medium']['body_min'])

    def e(self):
        return self.th['scale']['e_mult'] * self.g()

    def title_gap(self):
        return self.th['scale']['title_gap_step'] * self.g()

    def band(self):
        return self.th['scale']['bands'].get(self.mf.get('density'), {})

    def text_union_connected(self, shapes=None):
        """countable 文本墨迹框（各扩 g）是否单一连通块；返回 (bool, union_box|None)"""
        shapes = [s for s in (shapes if shapes is not None else self.countable) if s.is_text and s.text.strip()]
        if not shapes:
            return False, None
        pad = self.g()
        boxes = [s.ink.pad(pad) for s in shapes]
        parent = list(range(len(boxes)))

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                if boxes[i].intersects(boxes[j], eps=0.05):
                    parent[find(i)] = find(j)
        roots = {find(i) for i in range(len(boxes))}
        return len(roots) == 1, union_box([s.ink for s in shapes])

    # ---- 图片像素
    def image_pixels(self, si):
        """图片形状的可见像素（按 srcRect 裁剪后）PIL Image，缓存"""
        key = ('img', id(si))
        if key in self._cache:
            return self._cache[key]
        from PIL import Image
        import io
        try:
            im = Image.open(io.BytesIO(si.shape.image.blob)).convert('RGB')
            w, h = im.size
            cl, cr = si.shape.crop_left or 0, si.shape.crop_right or 0
            ct, cb = si.shape.crop_top or 0, si.shape.crop_bottom or 0
            im = im.crop((int(w * cl), int(h * ct), int(w * (1 - cr)), int(h * (1 - cb))))
        except Exception:
            im = None
        self._cache[key] = im
        return im

    def image_orig_size(self, si):
        from PIL import Image
        import io
        try:
            return Image.open(io.BytesIO(si.shape.image.blob)).size
        except Exception:
            return None


class DeckCtx:
    def __init__(self, prs, manifest, th, qa_dir, pptx_path):
        self.prs = prs
        self.manifest = manifest
        self.th = th
        self.qa = qa_dir
        self.pptx_path = pptx_path
        d = manifest['deck']
        self.deck = d
        self.page_w = prs.slide_width / EMU_PER_PT
        self.page_h = prs.slide_height / EMU_PER_PT
        self.margins = d['margins']
        self.palette_all = set()
        for k, v in d['palette'].items():
            for c in (v or []):
                self.palette_all.add(str(c).upper())
        self.accent = {str(c).upper() for c in d['palette'].get('accent', [])}
        pages_mf = {int(p['page']): p for p in manifest.get('pages', [])}
        self.pages = [PageCtx(self, i, s, pages_mf.get(i)) for i, s in enumerate(prs.slides, 1)]
        self._body_size = None

    def body_font_size(self):
        """deck 级正文字号：全 deck 内容页正文文本的 run 众数"""
        if self._body_size is None:
            c = Counter()
            for p in self.content_pages:
                c.update(p.body_size_counter())
            self._body_size = c.most_common(1)[0][0] if c else None
        return self._body_size

    @property
    def content_pages(self):
        return [p for p in self.pages if p.is_content]

    def notes_text(self, page):
        s = page.slide
        if s.has_notes_slide:
            return s.notes_slide.notes_text_frame.text or ''
        return ''


# ---------------------------------------------------------------------------- 检查注册

CHECKS = []


def check(cid, level, stage, desc, th_keys=()):
    def deco(fn):
        CHECKS.append({'id': cid, 'level': level, 'stage': stage, 'desc': desc, 'th': list(th_keys), 'fn': fn})
        return fn
    return deco


def F(page, cid, level, msg, **values):
    return Finding(page, cid, level, msg, values)


# ---------------------------------------------------------------------------- 1 schema

th_global = {}


def validate_schema(manifest, prs):
    """§1：结构与枚举；返回 findings（任一即整套不合格）"""
    out = []
    d = manifest.get('deck')
    if not isinstance(d, dict):
        return [F(None, 'schema', 'M', 'manifest 缺 deck 段')]
    for k in DECK_REQUIRED:
        if k not in d:
            out.append(F(None, 'schema', 'M', f'deck 缺字段 {k}'))
    if out:
        return out
    if d['tone'] not in ENUM['tone']:
        out.append(F(None, 'schema', 'M', f'deck.tone 非法 {d["tone"]!r}', allowed=sorted(ENUM['tone'])))
    for k in ('accent', 'secondary', 'tertiary', 'neutral'):
        if k not in d['palette']:
            out.append(F(None, 'schema', 'M', f'deck.palette 缺 {k}'))
    for k in ('l', 'r', 't', 'b'):
        if k not in d['margins']:
            out.append(F(None, 'schema', 'M', f'deck.margins 缺 {k}'))
    for a in ('title_anchor', 'pagenum_anchor'):
        for k in ('x', 'y', 'tol'):
            if k not in d[a]:
                out.append(F(None, 'schema', 'M', f'deck.{a} 缺 {k}'))
    for k in ('min', 'max'):
        if k not in d['title_size']:
            out.append(F(None, 'schema', 'M', f'deck.title_size 缺 {k}'))
    for k in ('corner', 'stroke'):
        if k not in d['shape_language']:
            out.append(F(None, 'schema', 'M', f'deck.shape_language 缺 {k}'))
    corner = d['shape_language'].get('corner')
    if isinstance(corner, dict) and set(corner) != {'small', 'medium', 'large'}:
        out.append(F(None, 'schema', 'M', 'deck.shape_language.corner 应为 {small, medium, large}', value=corner))
    if not (isinstance(d['fonts_latin'], list) and d['fonts_latin'] and all(f in d['fonts'] for f in d['fonts_latin'])):
        out.append(F(None, 'schema', 'M', 'deck.fonts_latin 必须是非空列表且包含在 deck.fonts 内', value=d['fonts_latin']))
    # 页面尺寸回填比对
    pw, ph = prs.slide_width / EMU_PER_PT, prs.slide_height / EMU_PER_PT
    if abs(d['page'].get('w', -1) - pw) > 0.5 or abs(d['page'].get('h', -1) - ph) > 0.5:
        out.append(F(None, 'schema', 'M', 'deck.page 与 pptx 页面尺寸不一致',
                     declared=d['page'], actual={'w': round(pw, 1), 'h': round(ph, 1)}))
    pages = manifest.get('pages')
    if not isinstance(pages, list) or not pages:
        return out + [F(None, 'schema', 'M', 'manifest 缺 pages')]
    n_slides = len(prs.slides)
    seen = set()
    for p in pages:
        if not isinstance(p, dict) or 'page' not in p:
            out.append(F(None, 'schema', 'M', 'pages 条目缺 page'))
            continue
        pg = int(p['page'])
        seen.add(pg)
        if 'type' not in p or p['type'] not in ENUM['type']:
            out.append(F(pg, 'schema', 'M', f'type 缺失或非法 {p.get("type")!r}', allowed=sorted(ENUM['type'])))
            continue
        if p['type'] != 'content':
            continue
        for k in CONTENT_REQUIRED:
            if k not in p:
                out.append(F(pg, 'schema', 'M', f'内容页缺字段 {k}'))
        for k in ('density', 'container', 'contrast', 'title_pos'):
            if k in p and p[k] not in ENUM[k]:
                out.append(F(pg, 'schema', 'M', f'{k} 非法 {p[k]!r}', allowed=sorted(ENUM[k])))
        if 'columns' in p and not COLUMNS_RE.fullmatch(str(p['columns'])):
            out.append(F(pg, 'schema', 'M', 'columns 格式非法，应为 "1"、"2:1"、"1:1:1" 这类列比', value=p['columns']))
        if 'hero' in p:
            out.append(F(pg, 'schema', 'M', 'hero 字段已改为 focus: {object, form, count}（先写这一页在讲谁，再选 form）'))
        fo = p.get('focus')
        if 'focus' in p:
            if not isinstance(fo, dict) or not all(k in fo for k in ('object', 'form', 'count')):
                out.append(F(pg, 'schema', 'M', 'focus 必须是 {object, form, count}', value=fo))
            else:
                if not (isinstance(fo['object'], str) and fo['object'].strip()):
                    out.append(F(pg, 'schema', 'M', 'focus.object 必须是非空字符串（这一页在讲谁）'))
                if fo['form'] not in ENUM['focus.form'] or fo['form'] == 'statement':
                    out.append(F(pg, 'schema', 'M', f'focus.form 非法 {fo["form"]!r}（statement 只用于 statement 页）', allowed=sorted(ENUM['focus.form'] - {'statement'})))
                if not (isinstance(fo['count'], int) and 0 <= fo['count'] <= 4):
                    out.append(F(pg, 'schema', 'M', 'focus.count 必须是 0–4 的整数', value=fo['count']))
        if 'visual_side' in p and p['visual_side'] not in ENUM['visual_side']:
            out.append(F(pg, 'schema', 'M', f'visual_side 非法 {p["visual_side"]!r}', allowed=sorted(ENUM['visual_side'])))
        if 'void' in p:
            out.append(F(pg, 'schema', 'M', 'void 字段已删除（留白不再声明），请改为 columns'))
        if 'g' in p and not (isinstance(p['g'], (int, float)) and p['g'] > 0):
            out.append(F(pg, 'schema', 'M', 'g 必须是正数（layout.py 算出后写入）', value=p['g']))
        if 'source_chars' in p and not (isinstance(p['source_chars'], int) and p['source_chars'] > 0):
            out.append(F(pg, 'schema', 'M', 'source_chars 必须是正整数（原稿分配给本页的汉字当量）', value=p['source_chars']))
        for k in ('subtitle', 'conclusion'):
            if k in p and not isinstance(p[k], bool):
                out.append(F(pg, 'schema', 'M', f'{k} 必须是布尔值'))
        if 'pattern' in p and not (isinstance(p['pattern'], str) and p['pattern'].strip()):
            out.append(F(pg, 'schema', 'M', 'pattern 必须是非空字符串'))
        if 'image' in p and p['image'] is not None:
            imgs = p['image'] if isinstance(p['image'], list) else [p['image']]
            for i, im in enumerate(imgs):
                if not isinstance(im, dict):
                    out.append(F(pg, 'schema', 'M', f'image[{i}] 不是映射'))
                    continue
                for k in IMAGE_REQUIRED:
                    if k not in im:
                        out.append(F(pg, 'schema', 'M', f'image[{i}] 缺字段 {k}'))
                if im.get('role') not in ENUM['image.role']:
                    out.append(F(pg, 'schema', 'M', f'image[{i}].role 非法 {im.get("role")!r}'))
                if str(im.get('source', '')).lower() not in ENUM['image.source']:
                    out.append(F(pg, 'schema', 'M', f'image[{i}].source 非法 {im.get("source")!r}', allowed=sorted(ENUM['image.source'])))
                if im.get('side') not in ENUM['image.side']:
                    out.append(F(pg, 'schema', 'M', f'image[{i}].side 非法 {im.get("side")!r}'))
                if not (isinstance(im.get('layout'), int) and 1 <= im['layout'] <= 7):
                    out.append(F(pg, 'schema', 'M', f'image[{i}].layout 非法 {im.get("layout")!r}，应为 1–7'))
        if isinstance(fo, dict) and fo.get('form') == 'image' and not p.get('image'):
            out.append(F(pg, 'schema', 'M', 'focus.form=image 但无 image 记录'))
    for i in range(1, n_slides + 1):
        if i not in seen:
            out.append(F(i, 'schema', 'M', f'pptx 第 {i} 页在 manifest 中缺失'))
    for pg in seen:
        if pg > n_slides:
            out.append(F(pg, 'schema', 'M', f'manifest 第 {pg} 页超出 pptx 页数 {n_slides}'))
    return out


# ---------------------------------------------------------------------------- 2 roles / C-01

def _max_run(s):
    rs = [r for r in s.runs if r.size and r.text.strip()]
    return max(rs, key=lambda r: r.size) if rs else None


@check('C-01', 'M', 'roles', '角色标记：形状名合法；内容页恰一个 title；hero:* 0–4 个且同页限定词 / 字号 / 字重 / 颜色一致；panel:region ≤ 1 且不与 card 同页', ('hierarchy.max_heroes',))
def c01_roles(ctx, page):
    out = [F(page.idx, 'C-01', 'M', e) for e in page.role_errors]
    if not page.is_content:
        return out
    for h in page.heroes:
        if h.qual == 'image' and not h.is_picture:
            out.append(F(page.idx, 'C-01', 'M', 'hero:image 形状不是图片', shape=h.name))
    sig = {}
    for h in page.heroes:
        if h.qual not in ('big-number', 'big-label'):
            continue
        r = _max_run(h)
        if r is not None:
            sig[f'{h.name}@{round(h.box.x)},{round(h.box.y)}'] = (r.size, r.bold, r.color)
    if len(set(sig.values())) > 1:
        out.append(F(page.idx, 'C-01', 'M', '同页多个 hero:* 的字号 / 字重 / 颜色必须一致（并列项同字号、同字重、同颜色）',
                     heroes={k: {'size': v[0], 'bold': v[1], 'color': v[2]} for k, v in sig.items()}))
    return out


# ---------------------------------------------------------------------------- 3 geometry

def _page_image_for_title(page):
    imgs = page.images
    return imgs[0] if imgs else None


@check('C-02', 'M', 'geometry', '标题锚点：按 title_pos 求期望位置比对 title 形状框；title_pos 与 image.layout/side 合法组合；特殊页 title 居中与字号',
       ('anchor.title_tol', 'anchor.mask_edge_frac', 'anchor.mask_edge_vtol', 'anchor.image_center_frac', 'anchor.cover_center_tol', 'scale.title_gap_step', 'cover_title_size'))
def c02_title_anchor(ctx, page):
    out = []
    th = ctx.th['anchor']
    t = page.title
    if not page.is_content:
        cts = ctx.th.get('cover_title_size', [45, 64])
        if page.type in ('cover', 'section') and t is not None:
            dcx = t.box.cx - ctx.page_w / 2
            if abs(dcx) > th['cover_center_tol']:
                out.append(F(page.idx, 'C-02', 'M', f'{page.type} 页 title 未水平居中', dcx=round(dcx, 1), tol=th['cover_center_tol']))
            sizes = [r.size for r in t.runs if r.size]
            if sizes and not (cts[0] <= max(sizes) <= cts[1]):
                out.append(F(page.idx, 'C-02', 'M', f'{page.type} 页 title 字号应在 {cts[0]}–{cts[1]}', size=max(sizes)))
        return out
    if t is None:
        return out
    anchor = ctx.deck['title_anchor']
    tol = anchor.get('tol', th['title_tol'])
    tp = page.mf.get('title_pos', 'base')
    img_mf = page.mf.get('image')
    img0 = (img_mf[0] if isinstance(img_mf, list) else img_mf) if img_mf else None
    # 合法组合
    if img0:
        layout, side = img0.get('layout'), img0.get('side')
        key = (layout, side if layout in (4, 5, 6) else None)
        legal = LEGAL_TITLE_POS.get(key)
        if legal is None:
            out.append(F(page.idx, 'C-02', 'M', 'image.layout/side 组合无合法 title_pos', layout=layout, side=side))
        elif tp not in legal:
            out.append(F(page.idx, 'C-02', 'M', 'title_pos 与 image.layout/side 组合不合法', title_pos=tp, layout=layout, side=side, allowed=sorted(legal)))
    elif tp != 'base':
        out.append(F(page.idx, 'C-02', 'M', '无图页 title_pos 必须为 base', title_pos=tp))
    img = _page_image_for_title(page)
    sg = page.title_gap()
    if tp == 'base':
        dx, dy = t.box.x - anchor['x'], t.box.y - anchor['y']
        if abs(dx) > tol or abs(dy) > tol:
            out.append(F(page.idx, 'C-02', 'M', 'title 左上角偏离基准锚点', dx=round(dx, 1), dy=round(dy, 1), tol=tol, expected=anchor, actual=t.box.r()))
    elif tp == 'below-image':
        if img is None:
            out.append(F(page.idx, 'C-02', 'M', 'title_pos=below-image 但页面无图片形状'))
        else:
            ex, ey = anchor['x'], img.box.bottom + sg
            dx, dy = t.box.x - ex, t.box.y - ey
            if abs(dx) > tol or abs(dy) > tol:
                out.append(F(page.idx, 'C-02', 'M', 'title 未落在图片下方 + 3g', dx=round(dx, 1), dy=round(dy, 1), tol=tol, expected={'x': ex, 'y': round(ey, 1)}, actual=t.box.r(), section_gap=round(sg, 1)))
    elif tp == 'right-of-image':
        if img is None:
            out.append(F(page.idx, 'C-02', 'M', 'title_pos=right-of-image 但页面无图片形状'))
        else:
            ex, ey = img.box.right + sg, anchor['y']
            dx, dy = t.box.x - ex, t.box.y - ey
            if abs(dx) > tol or abs(dy) > tol:
                out.append(F(page.idx, 'C-02', 'M', 'title 未落在图片右侧 + 3g', dx=round(dx, 1), dy=round(dy, 1), tol=tol, expected={'x': round(ex, 1), 'y': ey}, actual=t.box.r(), section_gap=round(sg, 1)))
    elif tp == 'mask-edge':
        m = page.masks[0] if page.masks else None
        if m is None:
            out.append(F(page.idx, 'C-02', 'M', 'title_pos=mask-edge 但页面无 mask'))
        else:
            edge = m.box.right if m.box.x <= ctx.page_w - m.box.right else m.box.x
            dcx, dcy = t.box.cx - edge, t.box.cy - m.box.cy
            if abs(dcx) > th['mask_edge_frac'] * ctx.page_w or abs(dcy) > th['mask_edge_vtol']:
                out.append(F(page.idx, 'C-02', 'M', 'title 中心未落在遮罩边缘', dcx=round(dcx, 1), dcy=round(dcy, 1),
                             tol_x=round(th['mask_edge_frac'] * ctx.page_w, 1), tol_y=th['mask_edge_vtol']))
    elif tp == 'image-center':
        if img is None:
            out.append(F(page.idx, 'C-02', 'M', 'title_pos=image-center 但页面无图片形状'))
        else:
            dcx, dcy = t.box.cx - img.box.cx, t.box.cy - img.box.cy
            if abs(dcx) > th['image_center_frac'] * img.box.w or abs(dcy) > th['image_center_frac'] * img.box.h:
                out.append(F(page.idx, 'C-02', 'M', 'title 中心未落在图片中心', dcx=round(dcx, 1), dcy=round(dcy, 1),
                             tol_x=round(th['image_center_frac'] * img.box.w, 1), tol_y=round(th['image_center_frac'] * img.box.h, 1)))
    return out


@check('C-03', 'W', 'geometry', '层级完整性：有大数字必有 label / heading；最大字号（title 除外）/ 正文 ≤ 3.5；heading ≥ 1.25 × 正文；大数字 ≤ 54；legend ≥ 18；hero:table / image 面积 ≥ 30% body [M]；hero:image 须 semantic [M]',
       ('hierarchy.max_ratio', 'hierarchy.heading_min_mult', 'hierarchy.number_max', 'legend_min_size', 'hero.area_frac'))
def c03_hierarchy(ctx, page):
    out = []
    if not page.is_content:
        return out
    hi = ctx.th['hierarchy']
    body_size = page.body_font_size()
    B = page.body_box
    ok = True
    numbers = [h for h in page.heroes if h.qual == 'big-number']
    if numbers and not (page.by_role.get('label') or page.by_role.get('heading')):
        ok = False
        out.append(F(page.idx, 'C-03', 'W', '有大数字但没有数字释义（label）或小标题级文字（heading）：Lonely Giant Number', numbers=[h.text for h in numbers]))
    biggest = None
    for s in page.foreground:
        if s.role in ('title', 'pagenum') or s.is_table:
            continue
        r = _max_run(s)
        if r is not None and (biggest is None or r.size > biggest[0]):
            biggest = (r.size, s.name)
    if biggest and body_size and biggest[0] / body_size > hi['max_ratio'] + 1e-6:
        ok = False
        out.append(F(page.idx, 'C-03', 'W', '最大字号 / 正文超过 3.5：层级断层', max_size=biggest[0], shape=biggest[1], body_size=body_size, ratio=round(biggest[0] / body_size, 2)))
    if body_size:
        small = {s.text[:12]: _max_run(s).size for s in page.by_role.get('heading', []) if _max_run(s) is not None and _max_run(s).size < hi['heading_min_mult'] * body_size - 1e-6}
        if small:
            ok = False
            out.append(F(page.idx, 'C-03', 'W', 'heading 字号不足 1.25 × 正文：小标题与正文同号', headings=small, body_size=body_size, required=round(hi['heading_min_mult'] * body_size, 1)))
    for h in numbers:
        r = _max_run(h)
        if r is not None and r.size > hi['number_max'] + 0.5:
            ok = False
            out.append(F(page.idx, 'C-03', 'W', '内容页大数字超过上限', size=r.size, max=hi['number_max'], text=h.text))
    lg = [(_max_run(s).size, s.text) for s in page.by_role.get('legend', []) if s.is_text and _max_run(s) is not None]
    if lg and min(v for v, _ in lg) < ctx.th.get('legend_min_size', 18) - 0.5:
        out.append(F(page.idx, 'C-03', 'W', '自绘序列名字号不足 18', sizes=[v for v, _ in lg]))
    for h in page.heroes:
        if h.qual in ('table', 'image') and B.area > 0 and h.box.area < ctx.th['hero']['area_frac'] * B.area:
            ok = False
            out.append(F(page.idx, 'C-03', 'M', f'hero:{h.qual} 面积不足 body 区 30%', hero_area=round(h.box.area), body_area=round(B.area), ratio=round(h.box.area / B.area, 3)))
        if h.qual == 'image':
            mf_img = page.mf.get('image')
            mf0 = (mf_img[0] if isinstance(mf_img, list) else mf_img) if mf_img else None
            if not mf0 or mf0.get('role') != 'semantic':
                ok = False
                out.append(F(page.idx, 'C-03', 'M', 'hero:image 的 manifest image.role 必须是 semantic', role=mf0.get('role') if mf0 else None))
    page.results['c03_ok'] = ok
    page.results['body_size'] = body_size
    return out


@check('C-04', 'M', 'geometry', '密度分档：汉字当量与元素数取较重者；> max_chars 拆页',
       ('density.light_chars', 'density.medium_chars', 'density.max_chars', 'density.light_elems', 'density.medium_elems'))
def c04_density(ctx, page):
    out = []
    if not page.is_content:
        return out
    th = ctx.th['density']
    chars, elems = page.density_chars(), page.element_count()
    if chars > th['max_chars']:
        out.append(F(page.idx, 'C-04', 'M', '单页字数超上限，必须拆页', chars=chars, limit=th['max_chars']))
    cb = 'light' if chars <= th['light_chars'] else 'medium' if chars <= th['medium_chars'] else 'heavy'
    eb = 'light' if elems <= th['light_elems'] else 'medium' if elems <= th['medium_elems'] else 'heavy'
    order = ['light', 'medium', 'heavy']
    band = order[max(order.index(cb), order.index(eb))]
    page.results['density'] = band
    page.results['density_values'] = {'chars': chars, 'elems': elems, 'chars_band': cb, 'elems_band': eb}
    return out


@check('C-09', 'M', 'geometry', '标题区下方：title / kicker 下沿到 body 第一个元素 ≥ 3g − e（只有下限；居中残余不限）', ('scale.title_gap_step', 'scale.e_mult'))
def c09_title_gap(ctx, page):
    out = []
    if not page.is_content or page.title is None:
        return out
    tb = page.title_bottom
    cb = page.content_box()
    if cb is None:
        return out
    expected = page.title_gap()
    gap = cb.y - tb
    e = page.e()
    if gap < expected - e:
        out.append(F(page.idx, 'C-09', 'M', '标题区到 body 第一个元素的距离小于 3g', gap=round(gap, 1), min=round(expected - e, 1), g=round(page.g(), 1), tol=round(e, 1)))
    page.results['title_gap'] = round(gap, 1)
    return out


def _in_band(v, lo, hi, tol):
    return (lo is None or v >= lo - tol) and (hi is None or v <= hi + tol)


@check('S-01', 'M', 'geometry', '尺度随密度：正文字号 ≥ 该档起始值（轻 18 / 中 16 / 重 14）；title 字号按档 [W]；hero:table / hero:image 高度 ≥ 80% body 高（图表不查）',
       ('scale.bands', 'scale.hero_block_h', 'type_scale_tol', 'title.heavy_size', 'title.other_size'))
def s01_scale(ctx, page):
    out = []
    if not page.is_content:
        return out
    band = page.band()
    tol = ctx.th['type_scale_tol']
    dens = page.mf.get('density')
    if page.body_size_counter():
        bs = page.body_font_size()
        if bs < band.get('body_min', 0) - tol:
            out.append(F(page.idx, 'S-01', 'M', f'{dens} 页正文字号低于该档起始值', body_size=bs, min=band.get('body_min')))
    t = page.title
    if t is not None and _max_run(t) is not None:
        lo, hi = ctx.th['title']['heavy_size'] if dens == 'heavy' else ctx.th['title']['other_size']
        ts = _max_run(t).size
        long_title = ts < lo and t.ink.w > ctx.th['title'].get('max_width_frac', 0.84) * t.box.w * ts / lo      # 长标题适当缩号：放到该档下限会超过 84% 标题宽
        if not (lo - tol <= ts <= hi + tol) and not long_title:
            out.append(F(page.idx, 'S-01', 'W', f'{dens} 页 title 字号不在该档区间（密度低 → 大标题也要放大）', size=ts, range=[lo, hi]))
    B = page.body_box
    need = ctx.th['scale']['hero_block_h'] * B.h
    for h in page.heroes:
        if h.qual in ('table', 'image') and h.box.h < need - 0.5:
            out.append(F(page.idx, 'S-01', 'M', f'hero:{h.qual} 高度不足 80% body 高', hero_h=round(h.box.h, 1), required=round(need, 1), body_h=round(B.h, 1)))
    return out


def _near_step(v, g, steps, tol):
    return any(abs(v - k * g) <= tol * k * g for k in steps)


@check('S-02', 'M', 'geometry', '间距：g 读 manifest（layout.py 写入），0.5 × 正文 ≤ g ≤ 1.5 × 正文；所有相邻间距 ∈ {g, 2g, 3g} ±20%；横向间距 = 2g / 3g',
       ('scale.g_min_mult', 'scale.g_max_mult', 'scale.steps', 'scale.step_tol', 'scale.column_gap_steps', 'scale.overlap'))
def s02_spacing(ctx, page):
    out = []
    if not page.is_content:
        return out
    sc = ctx.th['scale']
    gaps = page.gaps()
    g = page.g()
    bs = page.body_font_size() if page.body_size_counter() else float(page.band().get('body_min') or page.body_font_size())   # 无正文文本时按该档下限算 g 的上下限
    page.results['g'] = round(g, 1)
    page.results['body_size'] = bs
    page.results['gaps'] = sorted({round(x[0], 1) for x in gaps})
    lo, hi = sc['g_min_mult'] * bs, sc['g_max_mult'] * bs
    if g < lo - 0.5:
        out.append(F(page.idx, 'S-02', 'M', 'g 低于 0.5 × 正文字号：内容超载，拆页或降密度档', g=round(g, 1), body_size=bs, min=round(lo, 1)))
    elif g > hi + 0.5:
        out.append(F(page.idx, 'S-02', 'M', 'g 高于 1.5 × 正文字号：内容太小在撑开，放大元素', g=round(g, 1), body_size=bs, max=round(hi, 1)))
    bad = []
    rows = []
    for layer in [None] + page.containers:
        _align_edges([u for u in page.units() if u['layer'] is layer], sc['align_tol'], rows_out=rows, keep_units=True)
    for v, kind, a, b, *uab in gaps:
        alt = uab[2] if len(uab) > 2 else None
        uab = uab[:2] if len(uab) >= 2 and uab[0] is not None else []
        if a.role in ('title', 'kicker') and b.role not in ('title', 'kicker'):
            continue          # 标题区到 body 第一个元素的距离由 C-09（3g ± e）单独判定
        steps = sc['column_gap_steps'] if kind == 'h' else sc['steps']
        if kind == 'h' and uab:
            row = next((r for r in rows if any(u is uab[0] for u in r) and any(u is uab[1] for u in r)), None)
            if row is not None and len(row) >= 3:
                xs = sorted(u['box'].x for u in row)
                pitch = min(q - p_ for p_, q in zip(xs, xs[1:]))
                k = round((uab[1]['box'].x - uab[0]['box'].x) / pitch)
                if k >= 2:
                    v = v - (k - 1) * pitch                      # 并列组里的空位：扣掉整数个节距再判
        # 竖向有两种读法：两两距离、行到行（并排的格子里各自叠放时两两距离对，整行对齐时行到行对）；两种都不在档上才算违规
        if not _near_step(v, g, steps, sc['step_tol']) and not (alt is not None and _near_step(alt, g, steps, sc['step_tol'])):
            bad.append({'gap': round(v, 1), 'kind': kind, 'a': a.name, 'b': b.name, 'allowed': [round(k * g, 1) for k in steps]})
    if bad:
        out.append(F(page.idx, 'S-02', 'M', '间距不在 {g, 2g, 3g} ±20% 内（横向只允许 2g / 3g）', g=round(g, 1), bad=bad[:8], count=len(bad)))
    return out


@check('S-03', 'M', 'geometry', '内容块位置：横向左右沿 = 版心（贴页边的图片侧换成页边）± e；竖向填满 body ± e，或居中（上下残余相等 ± e，残余大小不限）',
       ('scale.e_mult',))
def s03_content_block(ctx, page):
    out = []
    if not page.is_content:
        return out
    cb = page.content_box()
    B, m = page.body_box, page.margin_box
    e = page.e()
    if cb is None:
        return [F(page.idx, 'S-03', 'M', '内容块为空')]
    L, R = m.x, m.right
    for s in page.foreground:
        if page.is_edge_image(s):                 # 满高边图：内容块的那一侧 = 图片内沿 − 分栏间距（2g–3g，取 2.5g ± e）
            if s.box.x <= 1:
                L = s.box.right + 2.5 * page.g()
            else:
                R = s.box.x - 2.5 * page.g()
        elif s.is_image_role:
            if s.box.x <= 1:
                L = 0.0
            if s.box.right >= ctx.page_w - 1:
                R = ctx.page_w
    region_l = any(r.box.x <= 1 and r.box.right < ctx.page_w - 1 for r in page.regions)
    region_r = any(r.box.right >= ctx.page_w - 1 and r.box.x > 1 for r in page.regions)
    d = {'left': cb.x - L, 'right': R - cb.right, 'top': cb.y - B.y, 'bottom': B.bottom - cb.bottom}
    center = cb.cy - B.cy
    page.results['content_box'] = cb.r()
    page.results['edges'] = {k: round(v, 1) for k, v in d.items()}
    page.results['edges_filled'] = abs(d['right']) <= e and abs(d['bottom']) <= e
    # 区域背景贴住的那一侧视为已填满：窄长区域里的文字与图收窄居中，只要求不越过版心边
    bad_l = d['left'] < -e if region_l else abs(d['left']) > e
    bad_r = d['right'] < -e if region_r else abs(d['right']) > e
    if bad_l or bad_r:
        out.append(F(page.idx, 'S-03', 'M', '内容块横向未填满版心', d_left=round(d['left'], 1), d_right=round(d['right'], 1), tol=round(e, 1), expected=[round(L, 1), round(R, 1)], content_box=cb.r()))
    filled = abs(d['top']) <= e and abs(d['bottom']) <= e
    centered = abs(d['top'] - d['bottom']) <= e
    page.results['vertical'] = 'filled' if filled else ('centered' if centered else 'neither')
    if not (filled or centered):
        out.append(F(page.idx, 'S-03', 'M', '内容块竖向既未填满 body 也未居中（上下残余不等）', d_top=round(d['top'], 1), d_bottom=round(d['bottom'], 1), d_center=round(center, 1),
                     tol=round(e, 1), body=B.r(), content_box=cb.r()))
    return out


def _max_empty_short(grid, cell):
    """最大空矩形短边（pt）"""
    if not grid:
        return 0.0
    rows, cols = len(grid), len(grid[0])
    best = 0
    height = [0] * cols
    for r in range(rows):
        for c in range(cols):
            height[c] = height[c] + 1 if grid[r][c] == 0 else 0
        stack = []
        for c in range(cols + 1):
            cur = height[c] if c < cols else 0
            start = c
            while stack and stack[-1][1] >= cur:
                sidx, h = stack.pop()
                best = max(best, min(h, c - sidx))
                start = sidx
            stack.append((start, cur))
    return best * cell


def _raster(region, boxes, cell):
    cols, rows = max(1, math.ceil(region.w / cell)), max(1, math.ceil(region.h / cell))
    grid = [[0] * cols for _ in range(rows)]
    for bx in boxes:
        b = bx.intersection(region)
        if b.area <= 0:
            continue
        c0, c1 = int((b.x - region.x) // cell), min(cols, math.ceil((b.right - region.x) / cell))
        r0, r1 = int((b.y - region.y) // cell), min(rows, math.ceil((b.bottom - region.y) / cell))
        for r in range(max(0, r0), r1):
            row = grid[r]
            for c in range(max(0, c0), c1):
                row[c] = 1
    return grid


@check('S-04', 'M', 'geometry', '容器贴内容：子项墨迹并集 = 容器内缩 p ± e（四边）；同页 p 一致且 ∈ [g, 3g]；容器内最大空矩形短边 ≤ 2p',
       ('scale.inset_min_mult', 'scale.inset_max_mult', 'scale.container_empty_mult', 'scale.e_mult', 'raster_cell'))
def s04_containers(ctx, page):
    out = []
    if not page.is_content or not page.containers:
        return out
    sc = ctx.th['scale']
    g, e = page.g(), page.e()
    cell = ctx.th['raster_cell']
    info = {}
    for c in page.containers:
        kids = page.children(c)
        if not kids:
            out.append(F(page.idx, 'S-04', 'M', '容器没有子项', container=c.name))
            continue
        inner = union_box([k.ink for k in kids])
        ins = {'left': inner.x - c.box.x, 'right': c.box.right - inner.right, 'top': inner.y - c.box.y, 'bottom': c.box.bottom - inner.bottom}
        info[c.name] = (ins, kids)
    if not info:
        return out
    p_by = {n: min(ins.values()) for n, (ins, _) in info.items()}
    p = min(p_by.values())
    page.results['p'] = round(p, 1)
    if max(p_by.values()) - p > e:
        out.append(F(page.idx, 'S-04', 'M', '各容器内缩 p 不一致', p={k: round(v, 1) for k, v in p_by.items()}, tol=round(e, 1)))
    if p < sc['inset_min_mult'] * g - e or p > sc['inset_max_mult'] * g + e:
        out.append(F(page.idx, 'S-04', 'M', 'p 不在 [g, 3g]', p=round(p, 1), g=round(g, 1), range=[round(sc['inset_min_mult'] * g, 1), round(sc['inset_max_mult'] * g, 1)]))
    for n, (ins, kids) in info.items():
        bad = {k: round(v, 1) for k, v in ins.items() if abs(v - p) > e}
        if bad:
            out.append(F(page.idx, 'S-04', 'M', '容器内容未贴四边（内缩 ≠ p ± e）', container=n, p=round(p, 1), tol=round(e, 1), insets={k: round(v, 1) for k, v in ins.items()}, bad=bad))
        c = next(x for x in page.containers if x.name == n)
        # 行带口径：子项横向占满容器内宽（右侧参差不算洞），竖向按墨迹
        bands = [Box(c.box.x, k.ink.y, c.box.w, k.ink.h) for k in kids]
        short = _max_empty_short(_raster(c.box, bands, cell), cell)
        limit = sc['container_empty_mult'] * p
        if short > limit + cell:
            out.append(F(page.idx, 'S-04', 'M', '容器内有短边 > 2p 的空矩形', container=n, empty_short_side=round(short, 1), max=round(limit, 1)))
    return out


@check('S-05', 'M', 'geometry', '无洞（行带口径）：content bbox 内栅格化，元素横向占满所在列、竖向按墨迹（分栏间距带、容器、并列组整体置 1），最大空矩形短边 ≤ 3g', ('scale.hole_mult', 'scale.overlap', 'raster_cell'))
def s05_holes(ctx, page):
    out = []
    if not page.is_content:
        return out
    cb = page.content_box()
    if cb is None:
        return out
    cell = ctx.th['raster_cell']
    cols = page.columns()
    # 行带口径：每个顶层元素横向占满它所在的列（右侧参差、居中小容器的两侧不算洞），竖向按墨迹；洞只剩行与行之间、列内上下残余
    boxes = [c.box for c in page.containers]
    cs_ids = {id(x) for x in page.content_shapes()}
    for u in page.units():
        if u['layer'] is not None or not any(id(x) in cs_ids for x in u['shapes']):
            continue
        boxes.append(u['ink'])
        for cb_, members in cols:
            if any(m is x for m in members for x in u['shapes']):
                boxes.append(Box(cb_.x, u['ink'].y, cb_.w, u['ink'].h))
    # 行（顶沿或垂直中心对齐的横排元素，如「小标题 | icon | 要点」）：行高由最高的那格决定，行的并集框整体置 1
    tol = ctx.th['scale']['align_tol']
    tops = [u for u in page.units() if u['layer'] is None and any(id(x) in cs_ids for x in u['shapes'])]
    regions = page.regions

    def region_of(u):
        return next((r for r in regions if r.box.contains(u['box'], tol)), None)
    for u in tops:
        mates = page.row_mates(u, tops, region_of)
        if len(mates) >= 2:
            boxes.append(union_box([v['ink'] for v in mates]))
    boxes += _parallel_rows([u for u in page.units() if u['layer'] is None], ctx.th['scale']['align_tol'])      # 并列组（含空位）整体不算洞
    for (a, _), (b, _) in zip(cols, cols[1:]):
        if b.x > a.right:
            boxes.append(Box(a.right, cb.y, b.x - a.right, cb.h))
    short = _max_empty_short(_raster(cb, boxes, cell), cell)
    limit = ctx.th['scale']['hole_mult'] * page.g()
    page.results['max_hole_short'] = round(short, 1)
    if short > limit + cell:
        out.append(F(page.idx, 'S-05', 'M', '内容块内有洞', hole_short_side=round(short, 1), max=round(limit, 1), g=round(page.g(), 1)))
    return out


def _cluster_1d(vals, tol):
    vals = sorted(vals)
    cl = []
    for v in vals:
        if cl and v - cl[-1][0] <= tol:
            cl[-1].append(v)
        else:
            cl.append([v])
    return [sum(c) / len(c) for c in cl]


def _parallel_rows(units, tol):
    """n ≥ 3 的并列组的并集框（时间线上没有事件的刻度是空位，不是洞）"""
    out = []
    _align_edges(units, tol, rows_out=out)
    return [union_box([units[k]['ink'] for k in row]) for row in out if len(row) >= 3]


def _align_edges(units, tol, pitch_tol=4.0, rows_out=None, keep_units=False):
    """并列组：同签名、顶沿或垂直中心对齐、等距（n ≥ 3，允许空位 = 最小间距的整数倍）或等宽（n = 2）横排的一组元素，
    整体只按首项左沿、末项右沿各计一条对齐线。返回 (lefts, rights)"""
    used, lefts, rights = set(), [], []
    order = sorted(range(len(units)), key=lambda i: units[i]['box'].x)
    for i in order:
        if i in used:
            continue
        a = units[i]
        row = [i]
        for j in order:
            if j == i or j in used:
                continue
            b = units[j]
            if b['sig'] == a['sig'] and (abs(b['box'].y - a['box'].y) <= tol or abs(b['box'].cy - a['box'].cy) <= tol) and b['box'].x > a['box'].x:
                row.append(j)
        ok = False
        if len(row) >= 3:
            xs = [units[k]['box'].x for k in row]
            d = [q - p for p, q in zip(xs, xs[1:])]
            base = min(d)
            ok = base > 0 and all(abs(v / base - round(v / base)) * base <= pitch_tol for v in d)
        elif len(row) == 2:
            ok = abs(units[row[0]]['box'].w - units[row[1]]['box'].w) <= tol
        if ok:
            used.update(row)
            if rows_out is not None:
                rows_out.append([units[k] for k in row] if keep_units else row)
            lefts.append(units[row[0]]['box'].x)
            rights.append(units[row[-1]]['box'].right)
        else:
            used.add(i)
            lefts.append(a['box'].x)
            rights.append(a['box'].right)
    return lefts, rights


@check('S-06', 'M', 'geometry', '对齐线（形状框）：顶层元素左沿 ≤ 3 个 x 值、右沿 ≤ 3 个（±2pt），title 左沿在其中；紧贴对算一个元素，等距横排的并列组只计首项左沿 / 末项右沿；容器与区域背景内以其内缘为基准另算 ≤ 3', ('scale.align_max', 'scale.align_tol'))
def s06_alignment(ctx, page):
    out = []
    if not page.is_content:
        return out
    sc = ctx.th['scale']
    tol, mx = sc['align_tol'], sc['align_max']

    def keep(u):
        return any(x.role not in ('arrow', 'divider', 'tag') for x in u['shapes']) and not any(page.is_edge_image(x) for x in u['shapes'])
    regions = page.regions

    def region_of(u):
        return next((r for r in regions if r.box.contains(u['box'], tol)), None)
    tops = [u for u in page.units() if u['layer'] is None and keep(u)]
    scopes = [('页面', [u for u in tops if region_of(u) is None])] + [(r.name, [u for u in tops if region_of(u) is r]) for r in regions]
    scopes += [(c.name, [u for u in page.units() if u['layer'] is c and keep(u)]) for c in page.containers]
    for name, us in scopes:
        if not us:
            continue
        # 居中叠放的一组（同一条中轴线、宽度不一，如窄长区域里的文字 + 数字 + 图）只算一条对齐线：取其中最宽者的左右沿
        axis = defaultdict(list)
        for u in us:
            axis[round(u['box'].cx / (2 * tol))].append(u)
        centered = [g_ for g_ in axis.values() if len(g_) >= 2 and len({round(u['box'].x) for u in g_}) >= 2]
        for g_ in centered:
            widest = max(g_, key=lambda u: u['box'].w)
            us = [u for u in us if not any(u is v for v in g_) or u is widest]
        ls, _ = _align_edges(us, tol)
        _, rs = _align_edges([u for u in us if not all(x.role in ('icon', 'legend') for x in u['shapes'])], tol)      # icon / 色块按自身尺寸，右沿不算对齐线
        lefts, rights = _cluster_1d(ls, tol), _cluster_1d(rs, tol)
        if name == '页面':
            page.results['align'] = {'left': [round(v, 1) for v in lefts], 'right': [round(v, 1) for v in rights]}
            t = page.title
            if t is not None and not any(abs(t.box.x - v) <= tol for v in lefts):
                out.append(F(page.idx, 'S-06', 'M', 'title 左沿不是左沿对齐线之一', title_left=round(t.box.x, 1), lefts=[round(v, 1) for v in lefts]))
        if len(lefts) > mx:
            out.append(F(page.idx, 'S-06', 'M', f'{name}：左沿超过 3 条对齐线', lefts=[round(v, 1) for v in lefts]))
        if len(rights) > mx:
            out.append(F(page.idx, 'S-06', 'M', f'{name}：右沿超过 3 条对齐线', rights=[round(v, 1) for v in rights]))
    return out


def _container_styles(page):
    styles, kinds = set(), {}
    bg = page.bg_color
    for c in page.cards:
        if c.fill and c.fill != bg:
            kinds[c.name] = 'fill'
        elif c.line:
            kinds[c.name] = 'stroke'
        else:
            kinds[c.name] = 'stroke' if not c.fill else 'fill'
    vals = list(kinds.values())
    selected = vals.count('fill') == 1 and vals.count('stroke') >= 2
    styles.update({'stroke'} if selected else set(vals))
    for p in page.by_role.get('panel', []):
        if p.qual == 'region':
            continue                                  # 区域背景是大容器，不计入中容器样式
        if p.fill and p.fill != bg:
            styles.add('fill')
    if page.by_role.get('divider'):
        styles.add('divider')
    return styles, kinds


@check('C-10', 'M', 'geometry', '容器样式：本页中容器样式集合 ≤ 1（恰一 fill + ≥2 stroke 的选中态记为 stroke；全 fill 的方案卡记为 fill）；卡片 2–6；描边框允许条件 [W]；panel:region 三边贴页边、颜色 ∈ tertiary ∪ secondary', ('cards', 'region'))
def c10_container(ctx, page):
    out = []
    if not page.is_content:
        return out
    styles, kinds = _container_styles(page)
    n = len(page.cards)
    vals = list(kinds.values())
    selected = vals.count('fill') == 1 and vals.count('stroke') >= 2
    page.results['container_set'] = sorted(styles)
    page.results['card_count'] = n
    if len(styles) > 1:
        out.append(F(page.idx, 'C-10', 'M', '本页容器样式超过一种', styles=sorted(styles), cards=kinds))
    cmin, cmax = ctx.th['cards']['min'], ctx.th['cards']['max']
    if n and not (cmin <= n <= cmax):
        out.append(F(page.idx, 'C-10', 'M', f'卡片数应在 {cmin}–{cmax}', cards=n))
    if 'stroke' in styles:
        multi = n >= 3 and all(
            sum(text_line_count(b.shape, ctx.th) for b in page.by_role.get('body', []) if c.box.contains(b.box)) >= 2
            for c in page.cards)
        duo = n == 2 and len({c.line for c in page.cards}) == 2          # 对比双框：两张描边颜色不同
        if not (multi or selected or duo):
            out.append(F(page.idx, 'C-10', 'W', '描边框不满足允许条件（≥3 张多行卡片，或恰一张填充选中态，或对比双框）', cards=n))
    rg = ctx.th['region']
    pal = ctx.deck['palette']
    allowed = {str(c).upper() for k in ('tertiary', 'secondary') for c in (pal.get(k) or [])}
    for r in page.regions:
        b = r.box
        edges = {'left': abs(b.x) <= rg['edge_tol'], 'right': abs(b.right - ctx.page_w) <= rg['edge_tol'],
                 'top': abs(b.y) <= rg['edge_tol'], 'bottom': abs(b.bottom - ctx.page_h) <= rg['edge_tol']}
        if sum(edges.values()) < rg['edges_min']:
            out.append(F(page.idx, 'C-10', 'M', 'panel:region 必须三边贴住页面边缘', box=b.r(), touching=[k for k, v in edges.items() if v]))
        if r.fill not in allowed:
            out.append(F(page.idx, 'C-10', 'M', 'panel:region 颜色须与 accent 同色系（palette.tertiary 或 secondary）', fill=r.fill))
        page.results['region_side'] = next((k for k in ('left', 'right', 'top', 'bottom') if not edges[{'left': 'right', 'right': 'left', 'top': 'bottom', 'bottom': 'top'}[k]]), 'full')
    return out


def _jaccard_2gram(a, b):
    def grams(s):
        s = re.sub(r'\s+', '', s)
        return {s[i:i + 2] for i in range(len(s) - 1)}
    ga, gb = grams(a), grams(b)
    if not ga or not gb:
        return 0.0
    return len(ga & gb) / len(ga | gb)


def _bare(t):
    return re.sub(r'[\s\W_]+', '', t or '')


@check('C-11', 'M', 'geometry', '副标题 / 结论行存在性与 manifest 一致；conclusion / heading / hero:big-label 与 title 的 2-gram Jaccard ≥ 0.6 视为复述（Title Echo）；title / kicker 不得是元标签',
       ('hierarchy.echo_jaccard', 'hierarchy.meta_labels'))
def c11_lines(ctx, page):
    out = []
    hi = ctx.th['hierarchy']
    meta = {_bare(m) for m in hi['meta_labels']}
    for s in page.by_role.get('title', []) + page.by_role.get('kicker', []):
        if _bare(s.text) in meta:
            out.append(F(page.idx, 'C-11', 'M', '元标签不上页面：只说明页面功能、不含信息的词不做 title / kicker', shape=s.name, text=s.text))
    if not page.is_content:
        return out
    has_k, has_c = bool(page.by_role.get('kicker')), bool(page.by_role.get('conclusion'))
    page.results['subtitle'], page.results['conclusion'] = has_k, has_c
    if has_k != bool(page.mf.get('subtitle')):
        out.append(F(page.idx, 'C-11', 'M', 'kicker 存在性与 manifest.subtitle 不符', declared=page.mf.get('subtitle'), actual=has_k))
    if has_c != bool(page.mf.get('conclusion')):
        out.append(F(page.idx, 'C-11', 'M', 'conclusion 存在性与 manifest.conclusion 不符', declared=page.mf.get('conclusion'), actual=has_c))
    if page.title is not None:
        cands = page.by_role.get('conclusion', []) + page.by_role.get('heading', []) + [h for h in page.heroes if h.qual == 'big-label']
        for s in cands:
            j = _jaccard_2gram(s.text, page.title.text)
            if j >= hi['echo_jaccard']:
                out.append(F(page.idx, 'C-11', 'M', '复述标题（Title Echo）', shape=s.name, text=s.text[:20], jaccard=round(j, 2), threshold=hi['echo_jaccard']))
    return out


_LATIN_RE = re.compile(r'[A-Za-z0-9]')
_CJK_CH = re.compile('[\u4e00-\u9fff]')


def _font_ok(f, allowed):
    return any(f == a or f.startswith(a + ' ') for a in allowed)


@check('C-13', 'M', 'geometry', '字体与字号：run 字体 ∈ fonts（前缀匹配字重名）；字号 ∈ type_scale ±0.5 且 ≥ min_size；内容页 title 28–40 且单行；本页字号种数 ≤ 5（不计 title / source / pagenum）；title 字间距 / 字号 ∈ [0.08, 0.12]；含数字 / 拉丁字符的 run 用英文字体；大数字粗体',
       ('type_scale_tol', 'max_sizes_per_page', 'title.spacing_frac'))
def c13_fonts(ctx, page):
    out = []
    d = ctx.deck
    tol, fonts, scale = ctx.th['type_scale_tol'], list(d['fonts']), d['type_scale']
    latin_fonts = list(d.get('fonts_latin') or [])
    sizes = set()
    bad_fonts, bad_sizes, small, bad_latin = set(), set(), set(), set()
    for s in page.shapes:
        for r in s.runs:
            if not r.text.strip():
                continue
            if not r.fonts:
                bad_fonts.add(f'{s.name}:继承')
            for f in r.fonts:
                if not _font_ok(f, fonts):
                    bad_fonts.add(f)
            if _LATIN_RE.search(r.text) and not (r.latin and _font_ok(r.latin, latin_fonts)):
                bad_latin.add(f'{s.name}:{r.text.strip()[:10]}→{r.latin}')
            if r.size is None:
                bad_sizes.add(f'{s.name}:继承')
                continue
            if s.role not in ('title', 'source', 'pagenum'):
                sizes.add(r.size)
            if not any(abs(r.size - t) <= tol for t in scale):
                bad_sizes.add(r.size)
            if r.size < d['min_size'] - tol:
                small.add(r.size)
    if bad_fonts:
        out.append(F(page.idx, 'C-13', 'M', '字体不在 deck.fonts', fonts=sorted(bad_fonts), allowed=sorted(fonts)))
    if bad_latin:
        out.append(F(page.idx, 'C-13', 'M', '数字 / 英文未用英文字体（a:latin 不在 deck.fonts_latin；生成后跑 scripts/postfix.py）', runs=sorted(bad_latin)[:8], count=len(bad_latin), allowed=latin_fonts))
    if bad_sizes:
        out.append(F(page.idx, 'C-13', 'M', '字号不在 type_scale', sizes=sorted(map(str, bad_sizes)), scale=scale))
    if small:
        out.append(F(page.idx, 'C-13', 'M', '字号低于 min_size', sizes=sorted(small), min=d['min_size']))
    t = page.title
    if page.is_content and t is not None:
        ts = [r.size for r in t.runs if r.size and r.text.strip()]
        if ts and not (d['title_size']['min'] <= max(ts) <= d['title_size']['max']):
            out.append(F(page.idx, 'C-13', 'M', 'title 字号超出 title_size 区间', size=max(ts), range=d['title_size']))
        n = text_line_count(t.shape, ctx.th)
        if n > 1:
            out.append(F(page.idx, 'C-13', 'M', 'title 必须单行：缩一档字号（不低于 title_size.min）', lines=n, text=t.text))
    if t is not None:
        lo, hi = ctx.th['title']['spacing_frac']
        bad = [{'text': r.text[:8], 'spc': r.spc, 'size': r.size} for r in t.runs
               if r.size and _CJK_CH.search(r.text) and not (lo - 1e-6 <= r.spc / r.size <= hi + 1e-6)]
        if bad:
            out.append(F(page.idx, 'C-13', 'M', '中文大标题字间距应为字号的 10% 左右（pptxgenjs charSpacing = 0.1 × fontSize）', runs=bad[:4], range=[lo, hi]))
    if page.is_content:
        heavy = [f'{s.role}:{s.text[:10]}' for s in page.by_role.get('heading', []) + page.by_role.get('conclusion', [])
                 if any(_CJK_CH.search(r.text) and (r.bold or any(f.endswith((' Bold', ' Heavy')) for f in r.fonts)) for r in s.runs)]
        if heavy:
            out.append(F(page.idx, 'C-13', 'W', '小标题 / 结论句用了大标题同款粗体：正文里除数字与个别关键词外不用 Bold，小标题用 Medium', shapes=heavy[:6]))
    for h in page.heroes:
        r = _max_run(h)
        if h.qual == 'big-number' and r is not None and not r.bold:
            out.append(F(page.idx, 'C-13', 'M', '放大强调的数字必须粗体', shape=h.name, text=h.text))
    if len(sizes) > ctx.th['max_sizes_per_page']:
        out.append(F(page.idx, 'C-13', 'M', '本页字号种数超过 5（不计 title / source / pagenum）', sizes=sorted(sizes)))
    return out


@check('C-14', 'M', 'geometry', '颜色：文字色、fill、line ∈ palette；accent 连续 run 汉字当量 ≤ accent_max_chars', ('accent_max_chars',))
def c14_colors(ctx, page):
    out = []
    pal = ctx.palette_all
    bad = set()
    for s in page.shapes:
        if s.is_picture:
            continue
        if s.fill and s.fill not in pal:
            bad.add(f'{s.name} fill {s.fill}')
        if s.line and s.line not in pal:
            bad.add(f'{s.name} line {s.line}')
        for c in _table_cell_fills(s):
            if c and c not in pal:
                bad.add(f'{s.name} cell {c}')
        for r in s.runs:
            if not r.text.strip():
                continue
            c = r.color or '000000'
            if c not in pal:
                bad.add(f'{s.name} text {c}')
    if bad:
        out.append(F(page.idx, 'C-14', 'M', '颜色不在 palette', colors=sorted(bad)))
    limit = ctx.deck.get('accent_max_chars', ctx.th['accent_max_chars'])
    for s in page.shapes:
        cur, cur_para = 0, None
        for r in s.runs:
            if r.color in ctx.accent and r.text.strip():
                cur = cur + cjk_equiv(r.text) if cur_para == r.para else cjk_equiv(r.text)
                cur_para = r.para
                if cur > limit:
                    out.append(F(page.idx, 'C-14', 'M', 'accent 色连续文字超限', shape=s.name, chars=cur, limit=limit))
                    break
            else:
                cur, cur_para = 0, None
    # 高亮色克制（01 Color，第六轮）：accent 不铺大面积；一页里 accent 文字不占多数
    ac = ctx.th.get('accent', {})
    page_area = ctx.page_w * ctx.page_h
    for s in page.shapes:
        if not s.is_picture and s.fill in ctx.accent and s.box.w * s.box.h > ac.get('fill_max_frac', 0.08) * page_area:
            out.append(F(page.idx, 'C-14', 'M', '高亮色铺了大面积（全屏 / 大色块）：饱和的高亮色不作背景色，改用低饱和的 tertiary / secondary 或中性色',
                         shape=s.name, frac=round(s.box.w * s.box.h / page_area, 2), limit=ac.get('fill_max_frac', 0.08)))
    tot = acc = 0
    for s in page.shapes:
        if s.role in ('title', 'pagenum', 'source'):
            continue
        for r in s.runs:
            n = len(r.text.strip())
            tot += n
            acc += n if r.color in ctx.accent else 0
    h = _hue(next(iter(ctx.accent), '000000'))
    lim = ac.get('text_max_frac_warm', 0.25) if (h < 70 or h > 300) else ac.get('text_max_frac', 0.35)
    if page.type == 'content' and tot >= 20 and acc / tot > lim:
        out.append(F(page.idx, 'C-14', 'W', '高亮色文字占比过高：高亮色留给本页结论里最重要的数字 / 词，论据数字、释义、小标题用深色或低饱和色（暖色高亮更要少用）',
                     accent_frac=round(acc / tot, 2), limit=lim))
    return out


def _hue(hex_):
    import colorsys
    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(r, g, b)[0] * 360


def _rel_lum(rgb):
    def ch(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = rgb
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def _hex_rgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


@check('C-15', 'M', 'geometry', '图上文字对比度：文本框与图片相交时，按遮罩合成后的像素亮度百分位算对比度；小字 ≥ 4.5，大字 ≥ 3.0',
       ('contrast.body', 'contrast.large', 'contrast.large_pt', 'contrast.large_bold_pt', 'contrast.percentile'))
def c15_contrast(ctx, page):
    out = []
    import numpy as np
    th = ctx.th['contrast']
    pct = th['percentile']
    for t in page.shapes:
        if not t.is_text or not t.text.strip():
            continue
        for im in page.images:
            if not t.box.intersects(im.box):
                continue
            pix = page.image_pixels(im)
            if pix is None:
                out.append(F(page.idx, 'C-15', 'M', '无法读取图片像素', image=im.name))
                continue
            inter = t.box.intersection(im.box)
            W, H = pix.size
            sx, sy = W / im.box.w, H / im.box.h
            x0, y0 = int((inter.x - im.box.x) * sx), int((inter.y - im.box.y) * sy)
            x1, y1 = max(x0 + 1, int((inter.right - im.box.x) * sx)), max(y0 + 1, int((inter.bottom - im.box.y) * sy))
            region = pix.crop((x0, y0, min(x1, W), min(y1, H)))
            region.thumbnail((256, 256))
            arr = np.asarray(region, dtype=float)
            for mk in page.masks:
                if mk.box.intersects(t.box) and mk.fill and not mk.fill.startswith('scheme') and len(mk.fill) == 6:
                    a = mk.fill_alpha
                    mc = np.array(_hex_rgb(mk.fill), dtype=float)
                    arr = mc * a + arr * (1 - a)
            lin = np.where(arr / 255.0 <= 0.03928, arr / 255.0 / 12.92, ((arr / 255.0 + 0.055) / 1.055) ** 2.4)
            L = 0.2126 * lin[..., 0] + 0.7152 * lin[..., 1] + 0.0722 * lin[..., 2]
            for r in t.runs:
                if not r.text.strip():
                    continue
                c = r.color or '000000'
                if len(c) != 6 or c.startswith('scheme'):
                    continue
                lt = _rel_lum(_hex_rgb(c))
                bg = float(np.percentile(L, pct)) if lt >= 0.5 else float(np.percentile(L, 100 - pct))
                contrast = (max(lt, bg) + 0.05) / (min(lt, bg) + 0.05)
                size = r.size or 18.0
                large = size >= th['large_pt'] or (r.bold and size >= th['large_bold_pt'])
                need = th['large'] if large else th['body']
                if contrast < need:
                    out.append(F(page.idx, 'C-15', 'M', '图上文字对比度不足', shape=t.name, image=im.name, contrast=round(contrast, 2),
                                 required=need, size=size, text_color=c, bg_L=round(bg, 3)))
                    break
    return out


@check('C-16', 'M', 'geometry', '页码：可无（≤ 20 页不放）；有则左上角与 pagenum_anchor 差 ≤ tol 且锚点在右下；> 20 页而无页码 [W]', ('pagenum_required_over',))
def c16_pagenum(ctx, page):
    out = []
    if not page.is_content:
        return out
    pn = page.by_role.get('pagenum', [])
    a = ctx.deck['pagenum_anchor']
    if not pn:
        if len(ctx.pages) > ctx.th.get('pagenum_required_over', 20):
            out.append(F(page.idx, 'C-16', 'W', '全 deck 超过 20 页，内容页应有页码（右下）', pages=len(ctx.pages)))
        return out
    if a['x'] < ctx.page_w / 2 or a['y'] < ctx.page_h / 2:
        out.append(F(page.idx, 'C-16', 'M', '页码锚点应在页面右下', anchor=a))
    dx, dy = pn[0].box.x - a['x'], pn[0].box.y - a['y']
    if abs(dx) > a['tol'] or abs(dy) > a['tol']:
        out.append(F(page.idx, 'C-16', 'M', '页码偏离锚点', dx=round(dx, 1), dy=round(dy, 1), tol=a['tol']))
    return out


# ---------------------------------------------------------------------------- 特殊页（§3.5）

@check('C-28', 'W', 'geometry', '一句话页（statement）：hero:big-label 字号按句长阶梯（≤7→72；8–14→60；15–24→48–54；25–40→40；>40 改要点页）；文字块宽 45–60% 页宽、高 20–35% 页高、整体居中',
       ('statement.ladder', 'statement.block_w', 'statement.block_h', 'statement.center_frac'))
def c28_statement(ctx, page):
    out = []
    if page.type not in ('statement', 'quote'):
        return out
    st = ctx.th['statement']
    hs = [h for h in page.heroes if h.qual == 'big-label']
    if len(hs) != 1:
        return [F(page.idx, 'C-28', 'W', '一句话页应恰有一个 hero:big-label（整页只有一个重点加粗放大）', count=len(hs))]
    h = hs[0]
    if page.title is not None:
        out.append(F(page.idx, 'C-28', 'W', '一句话页不带内容页标题：有实义的原标题降为辅助行（kicker），放在句子下方', title=page.title.text))
    n = cjk_equiv(h.text)
    r = _max_run(h)
    size = r.size if r is not None else None
    row = next((x for x in st['ladder'] if n <= x[0]), None)
    if row is None:
        out.append(F(page.idx, 'C-28', 'W', '超过 40 字不算一句话页，改要点页', chars=n))
    elif size is not None and not (row[1] - 0.5 <= size <= row[2] + 0.5):
        out.append(F(page.idx, 'C-28', 'W', '一句话页字号不在阶梯内（超过 7 个字就不能用最大号，否则视觉失衡）', chars=n, size=size, expected=[row[1], row[2]]))
    ink = h.ink
    fw, fh = ink.w / ctx.page_w, ink.h / ctx.page_h
    if not (st['block_w'][0] <= fw <= st['block_w'][1]) or not (st['block_h'][0] <= fh <= st['block_h'][1]):
        out.append(F(page.idx, 'C-28', 'W', '一句话文字块尺寸不在参照范围（宽 45–60%、高 20–35%）', w_frac=round(fw, 2), h_frac=round(fh, 2)))
    block = union_box([s.ink for s in page.foreground if s.is_text and s.role in ('hero', 'kicker', 'body')])
    dcx, dcy = block.cx - ctx.page_w / 2, block.cy - ctx.page_h / 2
    if abs(dcx) > st['center_frac'] * ctx.page_w or abs(dcy) > st['center_frac'] * ctx.page_h:
        out.append(F(page.idx, 'C-28', 'W', '一句话文字块未整体居中', dcx=round(dcx, 1), dcy=round(dcy, 1)))
    if r is not None and not r.bold:
        out.append(F(page.idx, 'C-28', 'W', '一句话页的重点句应加粗'))
    return out


@check('C-29', 'W', 'geometry', '目录页：有 image 且配图方式 ∈ {5, 6}；条目字号 ≥ 20', ('agenda.item_min_size', 'agenda.layouts'))
def c29_agenda(ctx, page):
    out = []
    if page.type != 'agenda':
        return out
    ag = ctx.th['agenda']
    imgs = _mf_images(page)
    if not imgs or not page.images:
        out.append(F(page.idx, 'C-29', 'W', '目录页也要排版：按配图方式 5 或 6 配图'))
    else:
        layout, side, note = _infer_layout(ctx, page, imgs[0].get('role'))
        if layout not in ag['layouts'] or layout != imgs[0].get('layout'):
            out.append(F(page.idx, 'C-29', 'W', '目录页配图方式不合法', declared=imgs[0].get('layout'), actual=layout, note=note))
    items = [_max_run(s).size for s in page.foreground if s.role in ('body', 'heading') and _max_run(s) is not None]
    if items and max(items) < ag['item_min_size'] - 0.5:
        out.append(F(page.idx, 'C-29', 'W', '目录条目字号不足 20', size=max(items)))
    if any(re.match(r'^\s*[一二三四五六七八九十]+[\s、.．]', s.text) for s in page.foreground if s.is_text):
        out.append(F(page.idx, 'C-29', 'W', '目录序号用英文字体加粗的数字（01–05）或圆形容器，不用「一 二 三」'))
    return out


# ---------------------------------------------------------------------------- deck 级

def _seq(ctx):
    return ctx.content_pages


@check('C-20', 'M', 'deck', '变化层：相邻页五项（pattern / focus.form / density / container / columns）≥ 2 不同（series 豁免）；5 页内同 pattern ≤ 2；相邻都有 card 时数量不同', ())
def c20_variation(ctx, _):
    out = []
    seq = _seq(ctx)
    keys = ['pattern', 'focus.form', 'density', 'container', 'columns']

    def val(p, k):
        return (p.mf.get('focus') or {}).get('form') if k == 'focus.form' else p.mf.get(k)
    for a, b in zip(seq, seq[1:]):
        same_series = a.mf.get('series') and a.mf.get('series') == b.mf.get('series')
        if same_series:
            continue
        diff = [k for k in keys if str(val(a, k)) != str(val(b, k))]
        if len(diff) < 2:
            out.append(F(b.idx, 'C-20', 'M', '相邻页 manifest 五项少于两项不同', prev=a.idx, differing=diff))
        ca, cb = a.results.get('card_count', 0), b.results.get('card_count', 0)
        if ca and cb and ca == cb:
            out.append(F(b.idx, 'C-20', 'M', '相邻页卡片数相同', prev=a.idx, cards=ca))
    for i in range(len(seq)):
        win = seq[i:i + 5]
        cnt = Counter(p.mf.get('pattern') for p in win if not p.mf.get('series'))
        for pat, n in cnt.items():
            if n > 2:
                out.append(F(win[0].idx, 'C-20', 'M', '连续 5 页内同一 pattern 超过 2 次', pattern=pat, count=n, pages=[p.idx for p in win if p.mf.get('pattern') == pat]))
                break
    return out


@check('C-21', 'W', 'deck', '系列页合计 ≤ 40% 内容页', ('deck.series_max',))
def c21_series(ctx, _):
    seq = _seq(ctx)
    if not seq:
        return []
    n = sum(1 for p in seq if p.mf.get('series'))
    if n / len(seq) > ctx.th['deck']['series_max']:
        return [F(None, 'C-21', 'W', '系列页占比过高', series=n, content=len(seq), ratio=round(n / len(seq), 2), max=ctx.th['deck']['series_max'])]
    return []


def _rhythm_seq(ctx):
    """节奏序列：content 页按 manifest.density；statement / section 页按 light 计入；cover / agenda / closing 不计。
    返回 (逐页序列 [(idx, density)], 滑动窗口序列)：窗口序列里一组连续的 series 页合成 1 项（取其中最重的档，页码取最后一页）"""
    order = ['light', 'medium', 'heavy']
    full = []
    for p in ctx.pages:
        if p.is_content:
            full.append((p.idx, p.mf.get('density'), p.mf.get('series')))
        elif p.type in ('statement', 'quote', 'section'):
            full.append((p.idx, 'light', None))
    win = []
    for idx, d, ser in full:
        if ser and win and win[-1][2] == ser:
            pd = win[-1][1]
            win[-1] = (idx, d if order.index(d) > order.index(pd) else pd, ser) if d in order and pd in order else (idx, pd, ser)
        else:
            win.append((idx, d, ser))
    return [(i, d) for i, d, _ in full], [(i, d) for i, d, _ in win]


@check('C-22', 'M', 'deck', '密度节奏（statement / section 按 light 计入；一组 series 页在滑动窗口里算 1 页）：连续 3 heavy [M]；连续 3 同档 [W]；heavy ≤ 40% [M]；heavy 后 2 页内有 light/medium [M]；前 3 / 末 2 含 light [W]', ('deck.heavy_max',))
def c22_rhythm(ctx, _):
    out = []
    full, win = _rhythm_seq(ctx)
    if not full:
        return out
    idx = [i for i, _ in win]
    d = [x for _, x in win]
    for i in range(len(d) - 2):
        if d[i] == d[i + 1] == d[i + 2]:
            lvl = 'M' if d[i] == 'heavy' else 'W'
            out.append(F(idx[i + 2], 'C-22', lvl, f'连续 3 页同为 {d[i]}', pages=[idx[i], idx[i + 1], idx[i + 2]]))
    fd = [x for _, x in full]                       # 占比按逐页序列算（series 不合并）
    heavy = fd.count('heavy')
    if heavy / len(fd) > ctx.th['deck']['heavy_max']:
        out.append(F(None, 'C-22', 'M', 'heavy 页占比超限', heavy=heavy, pages=len(fd), ratio=round(heavy / len(fd), 2), max=ctx.th['deck']['heavy_max']))
    for i, v in enumerate(d):
        if v == 'heavy' and i < len(d) - 1:
            nxt = d[i + 1:i + 3]
            if not any(x in ('light', 'medium') for x in nxt):
                out.append(F(idx[i], 'C-22', 'M', 'heavy 页之后 2 页内没有 light/medium', following=nxt))
    if 'light' not in d[:3]:
        out.append(F(None, 'C-22', 'W', '前 3 页内没有 light 页', first=d[:3]))
    if 'light' not in d[-2:]:
        out.append(F(None, 'C-22', 'W', '末 2 页内没有 light 页', last=d[-2:]))
    return out


@check('C-23', 'M', 'deck', '副标题页 ≤ 60% 且不连续 4 页；结论行连续 ≤ 2 且 ≤ 50%', ('deck.subtitle_max', 'deck.subtitle_run', 'deck.conclusion_run', 'deck.conclusion_max'))
def c23_lines_share(ctx, _):
    out = []
    seq = _seq(ctx)
    if not seq:
        return out
    th = ctx.th['deck']
    for key, mx, run, name in (('subtitle', th['subtitle_max'], th['subtitle_run'], '副标题'), ('conclusion', th['conclusion_max'], th['conclusion_run'], '结论行')):
        flags = [bool(p.mf.get(key)) for p in seq]
        n = sum(flags)
        if n / len(flags) > mx:
            out.append(F(None, 'C-23', 'M', f'{name}页占比超限', count=n, content=len(flags), ratio=round(n / len(flags), 2), max=mx))
        streak = 0
        limit = run - 1 if key == 'subtitle' else run     # 副标题：不得连续 4 页 → 最多 3；结论行：连续 ≤ 2
        for p, f in zip(seq, flags):
            streak = streak + 1 if f else 0
            if streak > limit:
                out.append(F(p.idx, 'C-23', 'M', f'{name}连续页数超限', streak=streak, limit=limit))
                break
    return out


@check('C-24', 'M', 'deck', '容器占比：容器样式集合为空的页 ≥ 1/3', ('deck.noframe_min',))
def c24_noframe(ctx, _):
    seq = _seq(ctx)
    if not seq:
        return []
    n = sum(1 for p in seq if not p.results.get('container_set'))
    if n / len(seq) < ctx.th['deck']['noframe_min']:
        return [F(None, 'C-24', 'M', '无容器页面不足 1/3', noframe=n, content=len(seq), ratio=round(n / len(seq), 2), min=ctx.th['deck']['noframe_min'])]
    return []


def _visual_side(page):
    imgs = _mf_images(page)
    if imgs:
        return imgs[0].get('side')
    return page.mf.get('visual_side')


@check('C-27', 'W', 'deck', '左右交替：相邻内容页的图表 / 图片 / 区域背景所在侧（image.side 或 visual_side）同为 left 或同为 right → 警告；series 豁免', ())
def c27_alternate(ctx, _):
    out = []
    seq = _seq(ctx)
    for a, b in zip(seq, seq[1:]):
        if a.mf.get('series') and a.mf.get('series') == b.mf.get('series'):
            continue
        sa, sb = _visual_side(a), _visual_side(b)
        if sa in ('left', 'right') and sa == sb:
            out.append(F(b.idx, 'C-27', 'W', '相邻页视觉元素在同一侧：上一页在左，下一页放右', prev=a.idx, side=sa))
    return out


@check('C-26', 'M', 'deck', '模板复制感综合：六项命中 ≥ 3 整套失败（第 5 项 = ≥ 80% 页内容块竖向填满而非居中）',
       ('composite.same_density', 'composite.both_lines', 'composite.same_container', 'composite.hierarchy_fail', 'composite.filled_edges', 'composite.pattern_share', 'composite.hits'))
def c26_composite(ctx, _):
    seq = _seq(ctx)
    if not seq:
        return []
    th = ctx.th['composite']
    n = len(seq)
    hits = {}
    dens = Counter(p.mf.get('density') for p in seq)
    hits['same_density'] = dens.most_common(1)[0][1] / n >= th['same_density']
    hits['both_lines'] = sum(1 for p in seq if p.mf.get('subtitle') and p.mf.get('conclusion')) / n >= th['both_lines']
    cont = Counter(tuple(p.results.get('container_set', [])) for p in seq)
    hits['same_container'] = cont.most_common(1)[0][1] / n >= th['same_container']
    hits['hierarchy_fail'] = sum(1 for p in seq if p.results.get('c03_ok') is False) / n >= th['hierarchy_fail']
    hits['filled_edges'] = sum(1 for p in seq if p.results.get('edges_filled')) / n >= th['filled_edges']
    pat = Counter(p.mf.get('pattern') for p in seq)
    hits['pattern_share'] = pat.most_common(1)[0][1] / n >= th['pattern_share']
    k = sum(hits.values())
    if k >= th['hits']:
        return [F(None, 'C-26', 'M', '模板复制感综合命中', hits=[h for h, v in hits.items() if v], count=k, threshold=th['hits'])]
    return []


# ---------------------------------------------------------------------------- 图片

def _mf_images(page):
    im = page.mf.get('image')
    if not im:
        return []
    return im if isinstance(im, list) else [im]


def _selected_files(ctx, page):
    d = os.path.join(ctx.qa, 'selected')
    if not os.path.isdir(d):
        return []
    pre = f'{page.idx:02d}'
    return [os.path.join(d, f) for f in sorted(os.listdir(d)) if f.startswith(pre) and (f[len(pre):len(pre) + 1] in ('.', '-'))]


@check('C-30', 'M', 'images', '选图记录：候选文件、viewed ≥ 8 且 = candidates_viewed、id 在候选内、reason ≥ 8 字、备注含 source 与 photographer、缩略图 ≤ 640px',
       ('images.viewed_min', 'images.reason_min', 'images.thumb_max_px'))
def c30_candidates(ctx, page):
    out = []
    if not page.is_content:
        return out
    th = ctx.th['images']
    imgs = [im for im in _mf_images(page) if str(im.get('source', '')).lower() in STOCK_SOURCES]
    if not imgs:
        return out
    path = os.path.join(ctx.qa, 'candidates', f'{page.idx:02d}.json')
    if not os.path.exists(path):
        return [F(page.idx, 'C-30', 'M', '候选记录文件不存在', path=path)]
    try:
        with open(path, encoding='utf-8') as f:
            rec = json.load(f)
    except Exception as e:
        return [F(page.idx, 'C-30', 'M', f'候选记录无法解析：{e}', path=path)]
    cands = rec.get('candidates', [])
    viewed = [c for c in cands if c.get('viewed') is True]
    ids = {str(c.get('id')) for c in cands}
    notes = ctx.notes_text(page)
    for im in imgs:
        cv = im.get('candidates_viewed')
        if len(viewed) < th['viewed_min']:
            out.append(F(page.idx, 'C-30', 'M', 'viewed=true 的候选不足', viewed=len(viewed), min=th['viewed_min']))
        if cv != len(viewed):
            out.append(F(page.idx, 'C-30', 'M', 'candidates_viewed 与记录不符', declared=cv, actual=len(viewed)))
        if str(im.get('id')) not in ids:
            out.append(F(page.idx, 'C-30', 'M', '选定图片不在候选列表内', id=im.get('id'), candidates=sorted(ids)[:20]))
        if len(str(im.get('reason', '')).strip()) < th['reason_min']:
            out.append(F(page.idx, 'C-30', 'M', 'reason 过短', reason=im.get('reason'), min=th['reason_min']))
        for key in ('source', 'photographer'):
            v = str(im.get(key, ''))
            if not v or v.lower() not in notes.lower():
                out.append(F(page.idx, 'C-30', 'M', f'页备注不含 {key}', value=v))
    from PIL import Image
    for c in viewed:
        tp = c.get('thumb')
        if not tp:
            out.append(F(page.idx, 'C-30', 'M', '候选缺 thumb 路径', id=c.get('id')))
            continue
        full = tp if os.path.isabs(tp) else os.path.join(os.path.dirname(os.path.abspath(ctx.pptx_path)), tp)
        if not os.path.exists(full):
            out.append(F(page.idx, 'C-30', 'M', '缩略图文件不存在', thumb=tp))
            continue
        try:
            w, h = Image.open(full).size
        except Exception:
            out.append(F(page.idx, 'C-30', 'M', '缩略图无法读取', thumb=tp))
            continue
        if max(w, h) > th['thumb_max_px']:
            out.append(F(page.idx, 'C-30', 'M', '缩略图长边超限（疑似原图）', thumb=tp, size=[w, h], max=th['thumb_max_px']))
    return out


@check('C-31', 'M', 'images', 'contact sheet 存在、晚于所有选中图、格数 ≥ 选中图数', ())
def c31_contact_sheet(ctx, _):
    pages = [p for p in ctx.content_pages if _mf_images(p)]
    if not pages:
        return []
    path = os.path.join(ctx.qa, 'contact_sheet.png')
    if not os.path.exists(path):
        return [F(None, 'C-31', 'M', 'contact sheet 不存在', path=path)]
    files = [f for p in pages for f in _selected_files(ctx, p)]
    n = sum(len(_mf_images(p)) for p in pages)
    out = []
    mt = os.path.getmtime(path)
    late = [f for f in files if os.path.getmtime(f) > mt]
    if late:
        out.append(F(None, 'C-31', 'M', 'contact sheet 早于选中图', newer=[os.path.basename(f) for f in late]))
    if not files:
        out.append(F(None, 'C-31', 'M', '_qa/selected/ 下没有选中图文件'))
    from PIL import Image
    from contact_sheet import CELL_W, CELL_H
    w, h = Image.open(path).size
    cells = (w // CELL_W) * (h // CELL_H)
    if cells < n:
        out.append(F(None, 'C-31', 'M', 'contact sheet 格数不足', cells=cells, images=n, size=[w, h]))
    return out


def _phash(im):
    import imagehash
    return imagehash.phash(im)


@check('C-32', 'M', 'images', '相邻内容页同角色图片：id 不同且 pHash 汉明距离 ≥ 12', ('images.phash_min',))
def c32_adjacent_images(ctx, _):
    out = []
    seq = ctx.content_pages
    mn = ctx.th['images']['phash_min']
    for a, b in zip(seq, seq[1:]):
        ra = defaultdict(list); rb = defaultdict(list)
        for im in _mf_images(a):
            ra[im.get('role')].append(str(im.get('id')))
        for im in _mf_images(b):
            rb[im.get('role')].append(str(im.get('id')))
        for role in set(ra) & set(rb):
            dup = set(ra[role]) & set(rb[role])
            if dup:
                out.append(F(b.idx, 'C-32', 'M', '相邻页同角色图片 id 相同', prev=a.idx, role=role, ids=sorted(dup)))
            sa = [s for s in a.images if s.image_qual == role]
            sb = [s for s in b.images if s.image_qual == role]
            for x in sa:
                px = a.image_pixels(x)
                if px is None:
                    continue
                hx = _phash(px)
                for y in sb:
                    py = b.image_pixels(y)
                    if py is None:
                        continue
                    dist = hx - _phash(py)
                    if dist < mn:
                        out.append(F(b.idx, 'C-32', 'M', '相邻页图片构图重复（pHash 距离不足）', prev=a.idx, role=role, distance=int(dist), min=mn))
    return out


@check('C-33', 'W', 'images', '摄影一致性：图库图 mean(R−B) 与 Lab L 均值的 z-score 绝对值 > 1.5 警告', ('images.temp_z',))
def c33_consistency(ctx, _):
    import numpy as np
    z = ctx.th['images']['temp_z']
    samples = []
    for p in ctx.content_pages:
        stock = {im.get('role') for im in _mf_images(p) if str(im.get('source', '')).lower() in STOCK_SOURCES}
        for s in p.images:
            if s.image_qual not in stock:
                continue
            px = p.image_pixels(s)
            if px is None:
                continue
            small = px.copy(); small.thumbnail((128, 128))
            arr = np.asarray(small, dtype=float)
            temp = float((arr[..., 0] - arr[..., 2]).mean())
            lin = np.where(arr / 255.0 <= 0.03928, arr / 255.0 / 12.92, ((arr / 255.0 + 0.055) / 1.055) ** 2.4)
            Y = 0.2126 * lin[..., 0] + 0.7152 * lin[..., 1] + 0.0722 * lin[..., 2]
            fy = np.where(Y > 0.008856, np.cbrt(Y), 7.787 * Y + 16 / 116)
            Lstar = float((116 * fy - 16).mean())
            samples.append((p.idx, s.name, temp, Lstar))
    if len(samples) < 3:
        return []
    out = []
    for k, label in ((2, '色温(R−B)'), (3, 'Lab L')):
        vals = np.array([s[k] for s in samples])
        sd = vals.std()
        if sd == 0:
            continue
        zs = (vals - vals.mean()) / sd
        for s, zz in zip(samples, zs):
            if abs(zz) > z:
                out.append(F(s[0], 'C-33', 'W', f'{label} 离群', image=s[1], value=round(float(s[k]), 1), z=round(float(zz), 2), threshold=z))
    return out


@check('C-35', 'M', 'geometry', '图片不变形：形状宽高比与（裁剪后）像素宽高比相对差 ≤ 2%', ('image_aspect_tol',))
def c35_aspect(ctx, page):
    out = []
    tol = ctx.th['image_aspect_tol']
    for s in page.images:
        if not s.is_picture:
            continue
        sz = page.image_orig_size(s)
        if sz is None or s.box.h <= 0:
            continue
        sh = s.shape
        cl, cr = sh.crop_left or 0, sh.crop_right or 0
        ct, cb = sh.crop_top or 0, sh.crop_bottom or 0
        pw, ph = sz[0] * (1 - cl - cr), sz[1] * (1 - ct - cb)
        if pw <= 0 or ph <= 0:
            continue
        pa, sa = pw / ph, s.box.w / s.box.h
        diff = abs(sa - pa) / pa
        if diff > tol:
            out.append(F(page.idx, 'C-35', 'M', '图片被拉伸变形', image=s.name, shape_aspect=round(sa, 3), pixel_aspect=round(pa, 3), diff=round(diff, 3), tol=tol,
                         cropped=bool(cl or cr or ct or cb)))
    return out


@check('C-36', 'M', 'geometry', '段落上限：单段 ≤ 100 汉字当量且 ≤ 4 行；单个文本框 ≤ 200 汉字当量（title / pagenum / source / 表格除外）',
       ('text.para_max_chars', 'text.para_max_lines', 'text.box_max_chars'))
def c36_paragraph(ctx, page):
    out = []
    if not page.is_content:
        return out
    th = ctx.th['text']
    from inkbox import _para_lines, line_width, DEFAULT_TH, _emu_attr
    ink = ctx.th.get('ink', DEFAULT_TH)
    for s in page.foreground:
        if not s.is_text or s.is_table or s.role in ('title', 'pagenum', 'source'):
            continue
        total = cjk_equiv(s.text)
        if total > th['box_max_chars']:
            out.append(F(page.idx, 'C-36', 'M', '文本框字数超上限', shape=s.name, chars=total, max=th['box_max_chars']))
        body_pr = s.shape.text_frame._txBody.bodyPr
        inner_w = max(s.box.w - _emu_attr(body_pr, 'lIns', ink['inset']) - _emu_attr(body_pr, 'rIns', ink['inset']), 1.0)
        for pi, p in enumerate(s.shape.text_frame.paragraphs):
            text = ''.join(r.text for r in s.runs if r.para == pi)
            chars = cjk_equiv(text)
            n = 0
            for line in _para_lines(p._p, 18.0):
                lw = sum(line_width(t, sz, ink, sp) for t, sz, sp in line)
                n += max(1, math.ceil(lw / inner_w))
            if chars > th['para_max_chars'] or n > th['para_max_lines']:
                out.append(F(page.idx, 'C-36', 'M', '段落超上限', shape=s.name, paragraph=pi, chars=chars, lines=n,
                             max_chars=th['para_max_chars'], max_lines=th['para_max_lines']))
    return out


@check('C-37', 'M', 'geometry', '字数锁定：本页正文汉字当量（C-04 口径）≤ source_mult × manifest.source_chars；排版不得改内容', ('text.source_mult',))
def c37_source_chars(ctx, page):
    out = []
    if not page.is_content:
        return out
    src = page.mf.get('source_chars')
    if not isinstance(src, int) or src <= 0:
        return out
    free = ctx.th['hierarchy']['heading_free_chars']
    # 提炼的小标题（heading / tag，≤ 10 字）不计入字数（01 字数锁定）
    chars = page.density_chars() - sum(cjk_equiv(s.text) for s in page.foreground if s.role in ('heading', 'tag') and cjk_equiv(s.text) <= free)
    limit = ctx.th['text']['source_mult'] * src
    page.results['source_chars'] = src
    if chars > limit:
        out.append(F(page.idx, 'C-37', 'M', '页面字数超过原稿分配的 1.1 倍：为了排版加了字', chars=chars, source_chars=src, max=round(limit)))
    return out


@check('C-34', 'M', 'images', '分辨率：full-bleed 原图宽 ≥ 1920；其它 ≥ 2 × 形状框宽(pt) × 96/72', ('images.fullbleed_min_px', 'images.small_mult'))
def c34_resolution(ctx, page):
    out = []
    th = ctx.th['images']
    for s in page.images:
        sz = page.image_orig_size(s)
        if sz is None:
            out.append(F(page.idx, 'C-34', 'M', '无法读取原图尺寸', image=s.name))
            continue
        if s.image_qual == 'full-bleed':
            need = th['fullbleed_min_px']
        else:
            need = th['small_mult'] * s.box.w * 96 / 72
        if sz[0] < need:
            out.append(F(page.idx, 'C-34', 'M', '原图宽度不足', image=s.name, width=sz[0], required=round(need)))
    return out


# ---------------------------------------------------------------------------- crosscheck §7

def _infer_layout(ctx, page, role):
    """§7：由形状框比例反算 image.layout 与 side；返回 (layout|None, side|None, note)"""
    tol = IMAGE_LAYOUT_TOL
    W, H = ctx.page_w, ctx.page_h
    shapes = [s for s in page.images if s.image_qual == role]
    if not shapes:
        return None, None, '无该角色图片形状'
    if len(shapes) >= 2 or role == 'small':
        ub = union_box([s.box for s in shapes])
        side = 'right' if ub.x >= 0.5 * W else ('top' if ub.cy < 0.5 * H else 'bottom')
        return 7, side, f'{len(shapes)} 张小图'
    b = shapes[0].box

    def near(v, target, base):
        return abs(v - target) <= tol * base
    full_w, full_h = near(b.w, W, W), near(b.h, H, H)
    if full_w and full_h:
        masks = page.masks
        for m in masks:
            if near(m.box.w, 0.74 * W, W) and near(m.box.h, H, H):
                return 3, 'full', '全屏 + 74% 遮罩'
            if near(m.box.w, W, W) and near(m.box.h, H, H):
                if m.fill and len(m.fill) == 6:
                    return (1 if _rel_lum(_hex_rgb(m.fill)) < 0.5 else 2), 'full', '全屏 + 全页遮罩'
                return None, 'full', '全页遮罩颜色不可读'
        return None, 'full', '全屏但无可识别遮罩'
    if full_w and near(b.h, 0.67 * H, H):
        return 4, ('top' if b.y <= tol * H else 'bottom'), '67% 高'
    for frac, n in ((0.22, 5), (0.50, 6)):
        if near(b.w, frac * W, W) or (n == 5 and 0.22 * W <= b.w <= 0.30 * W + tol * W):      # 方式 5：22–30% 宽
            # 方式 5 / 6 必须满页高并贴住页面上、下和一侧边缘（浮在页面中间的小图不属于任何配图方式）
            edge = b.x <= 1 or b.right >= W - 1
            if not (b.y <= 1 and b.bottom >= H - 1 and edge):
                return None, None, f'{int(frac * 100)}% 宽但未满页高 / 未贴边（Floating Image）：形状框 {b.r()}'
            return n, ('left' if b.cx < W / 2 else 'right'), f'{int(frac * 100)}% 宽，满高贴边'
    return None, None, f'形状框 {b.r()} 不匹配任何配图方式'


@check('C-40', 'M', 'crosscheck', 'manifest 声明 vs 几何反算：type / density / hero / container / subtitle / conclusion / image.role-layout-side', ())
def c40_crosscheck(ctx, page):
    out = []
    mf, res = page.mf, page.results

    def mismatch(field, declared, actual, **extra):
        out.append(F(page.idx, 'C-40', 'M', f'{field} 声明与反算不符', field=field, declared=declared, actual=actual, **extra))
    if page.type == 'cover' and page.idx != 1:
        pass  # 页码 1 = cover 允许；其余以 manifest 为准
    if page.type == 'section':
        if page.hero is not None:
            mismatch('type', 'section', 'has hero', shape=page.hero.name)
        nb = len(page.by_role.get('body', []))
        if nb >= 3:
            mismatch('type', 'section', f'body×{nb}')
    if not page.is_content:
        return out
    if 'density' in res and res['density'] != mf.get('density'):
        mismatch('density', mf.get('density'), res['density'], **res.get('density_values', {}))
    fo = mf.get('focus') or {}
    form, cnt = fo.get('form'), fo.get('count')
    quals = Counter(h.qual for h in page.heroes)
    n_num = quals.get('big-number', 0)
    n_head = len(page.by_role.get('heading', [])) + sum(1 for t in page.by_role.get('tag', []) if t.text.strip())     # 胶囊 / 圆形容器里的小标题也算
    actual = {'hero': dict(quals), 'heading': n_head, 'arrow': len(page.by_role.get('arrow', [])), 'tag': len(page.by_role.get('tag', []))}
    if form == 'numbers' and n_num != cnt:
        mismatch('focus', fo, actual, note='form=numbers：hero:big-number 个数须等于 count')
    elif form == 'headings' and n_head < (cnt or 0):
        mismatch('focus', fo, actual, note='form=headings：heading 个数须 ≥ count')
    elif form in ('chart', 'table', 'image') and quals.get(form, 0) != 1:
        mismatch('focus', fo, actual, note=f'form={form}：须恰有一个 hero:{form}')
    elif form == 'timeline' and (n_num or actual['arrow'] < 1 or actual['tag'] < 3):
        mismatch('focus', fo, actual, note='form=timeline：轴线 arrow ≥ 1、节点 tag ≥ 3，且不得挑一个时间做大数字')
    if 'container_set' in res:
        cs = res['container_set']
        actual_c = 'none' if not cs else (cs[0] if len(cs) == 1 else '+'.join(cs))
        if actual_c != mf.get('container'):
            mismatch('container', mf.get('container'), actual_c)
    for k in ('subtitle', 'conclusion'):
        if k in res and res[k] != bool(mf.get(k)):
            mismatch(k, mf.get(k), res[k])
    imgs = _mf_images(page)
    shape_roles = Counter(s.image_qual for s in page.images)
    mf_roles = Counter(im.get('role') for im in imgs)
    if shape_roles != mf_roles:
        mismatch('image.role', dict(mf_roles), dict(shape_roles))
    for im in imgs:
        layout, side, note = _infer_layout(ctx, page, im.get('role'))
        if layout != im.get('layout') or (side is not None and side != im.get('side')):
            mismatch('image.layout/side', {'layout': im.get('layout'), 'side': im.get('side')}, {'layout': layout, 'side': side}, note=note)
    return out


# ---------------------------------------------------------------------------- 主流程

STAGE_ORDER = ['schema', 'roles', 'geometry', 'deck', 'images', 'crosscheck', 'composite']


def load_thresholds(path):
    import yaml
    with open(path, encoding='utf-8') as f:
        return yaml.safe_load(f)


def run(pptx_path, manifest_path, th_path, qa_dir, json_out):
    import yaml
    from pptx import Presentation
    th = load_thresholds(th_path)
    th_global.update(th)
    with open(manifest_path, encoding='utf-8') as f:
        manifest = yaml.safe_load(f)
    prs = Presentation(pptx_path)
    findings = []
    schema = validate_schema(manifest, prs)
    if schema:
        findings.extend(schema)
        return finish(findings, json_out, stopped='schema')
    ctx = DeckCtx(prs, manifest, th, qa_dir, pptx_path)
    by_stage = defaultdict(list)
    for c in CHECKS:
        by_stage[c['stage']].append(c)
    page_ok = {}
    for page in ctx.pages:
        fs = c01_roles(ctx, page)
        findings.extend(fs)
        page_ok[page.idx] = not fs
    per_page_stages = ['geometry', 'images', 'crosscheck']
    for page in ctx.pages:
        if not page_ok[page.idx]:
            continue
        for c in by_stage['geometry']:
            if c['id'] == 'C-01':
                continue
            findings.extend(_safe(c, ctx, page))
    for c in by_stage['deck']:
        if c['id'] == 'C-26':
            continue
        findings.extend(_safe(c, ctx, None))
    for c in by_stage['images']:
        if c['id'] in ('C-31', 'C-32', 'C-33'):
            findings.extend(_safe(c, ctx, None))
        else:
            for page in ctx.pages:
                if page_ok[page.idx]:
                    findings.extend(_safe(c, ctx, page))
    for page in ctx.pages:
        if page_ok[page.idx]:
            findings.extend(_safe(next(c for c in CHECKS if c['id'] == 'C-40'), ctx, page))
    findings.extend(_safe(next(c for c in CHECKS if c['id'] == 'C-26'), ctx, None))
    pages_summary = {p.idx: {'type': p.type, **p.results} for p in ctx.pages}
    return finish(findings, json_out, pages=pages_summary)


def _safe(c, ctx, page):
    try:
        return c['fn'](ctx, page) or []
    except Exception as e:  # 校验器自身错误也要浮出来，不能静默通过
        import traceback
        return [F(page.idx if page else None, c['id'], 'M', f'校验器异常：{type(e).__name__}: {e}', trace=traceback.format_exc().splitlines()[-3:])]


def finish(findings, json_out, stopped=None, pages=None):
    m = [f for f in findings if f.level == 'M']
    w = [f for f in findings if f.level == 'W']
    report = {'stopped_at': stopped, 'summary': {'M': len(m), 'W': len(w), 'pass': not m},
              'findings': [asdict(f) for f in sorted(findings, key=lambda f: (f.page or 0, f.check))],
              'pages': pages or {}}
    if json_out:
        os.makedirs(os.path.dirname(json_out) or '.', exist_ok=True)
        with open(json_out, 'w', encoding='utf-8') as f:
            json.dump(report, f, ensure_ascii=False, indent=2, default=str)
    by_page = defaultdict(list)
    for f in findings:
        by_page[f.page].append(f)
    for pg in sorted(by_page, key=lambda x: (x is None, x or 0)):
        print(f'--- {"deck" if pg is None else f"page {pg:02d}"}')
        for f in by_page[pg]:
            vals = ' '.join(f'{k}={v}' for k, v in f.values.items() if k != 'trace')
            print(f'  [{f.level}] {f.check} {f.message}  {vals}')
    if stopped:
        print(f'\n在 {stopped} 阶段停止')
    print(f'\nM 失败 {len(m)}，W 警告 {len(w)} → {"通过" if not m else "不合格"}')
    if json_out:
        print(f'JSON: {json_out}')
    return 0 if not m else 1


def list_checks(th_path):
    th = load_thresholds(th_path) if os.path.exists(th_path) else {}

    def get(key):
        cur = th
        for part in key.split('.'):
            cur = cur.get(part, {}) if isinstance(cur, dict) else {}
        return json.dumps(cur, ensure_ascii=False) if isinstance(cur, dict) else cur
    print(f'{"ID":6} {"级别":3} {"阶段":10} 说明 / 阈值来源')
    print('schema M   schema     manifest 结构与枚举（§1），失败即停')
    for c in CHECKS:
        print(f'{c["id"]:6} {c["level"]:3} {c["stage"]:10} {c["desc"]}')
        if c['th']:
            print(f'{"":21} thresholds.yaml: ' + ', '.join(f'{k}={get(k)}' for k in c['th']))
        else:
            print(f'{"":21} 阈值：无数值阈值或来自 manifest.deck')
        if c['id'] == 'C-40':
            print(f'{"":21} 代码常量：IMAGE_LAYOUT_TOL={IMAGE_LAYOUT_TOL}（§7 比例判定容差，§8 未列）')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('pptx', nargs='?')
    ap.add_argument('manifest', nargs='?')
    ap.add_argument('--thresholds', default=os.path.join(HERE, '..', 'thresholds.yaml'))
    ap.add_argument('--qa', default=None, help='_qa 目录，默认 pptx 同目录下的 _qa')
    ap.add_argument('--json', default=None, help='JSON 输出路径，默认 <qa>/validate.json')
    ap.add_argument('--list', action='store_true')
    a = ap.parse_args()
    if a.list:
        list_checks(a.thresholds)
        return 0
    if not a.pptx or not a.manifest:
        ap.print_help()
        return 2
    qa = a.qa or os.path.join(os.path.dirname(os.path.abspath(a.pptx)), '_qa')
    json_out = a.json or os.path.join(qa, 'validate.json')
    return run(a.pptx, a.manifest, a.thresholds, qa, json_out)


if __name__ == '__main__':
    sys.exit(main())
