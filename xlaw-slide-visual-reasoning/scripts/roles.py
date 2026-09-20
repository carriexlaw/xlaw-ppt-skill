"""
roles.py — 形状角色标记校验（对应 skill 00「形状角色标记」）

用法：
    python roles.py deck.pptx [deck.manifest.yaml]   # 打印每页角色清单 + 违规；给 manifest 则按页类型判定
    from roles import classify_slide       # 在 validate_design.py 里调用

classify_slide(slide, page_type) 返回 (roles, errors)：
    roles  : dict  role -> [ (shape, qualifier) ]   供其它检查按角色取形状
    errors : list  本页违规描述，空列表即通过
"""
import sys
from collections import defaultdict

ROLES = {
    'title', 'kicker', 'conclusion', 'body', 'heading', 'label', 'legend',
    'hero', 'card', 'panel', 'tag', 'divider', 'mask',
    'image', 'icon', 'chart', 'table', 'arrow',
    'source', 'pagenum', 'bg', 'deco',
}
HERO_QUALIFIERS = {'big-number', 'big-label', 'chart', 'table', 'image'}
IMAGE_QUALIFIERS = {'full-bleed', 'atmosphere', 'semantic', 'small'}
PANEL_QUALIFIERS = {'', 'region'}           # panel:region = 贴边的区域背景
CONTENT_TYPES = {'content'}                 # 需要恰一个 title 的页类型；hero:* 0–MAX_HEROES 个（第一层级可以是一组）
SPECIAL_TYPES = {'cover', 'agenda', 'section', 'closing', 'statement', 'quote'}   # quote = statement 的旧名
MAX_HEROES = 4


def parse_role(name: str):
    """'hero:big-number' -> ('hero', 'big-number'); 'title' -> ('title', '')"""
    role, _, qual = (name or '').partition(':')
    return role, qual


def classify_slide(slide, page_type: str = 'content', max_heroes: int = MAX_HEROES):
    """内容页：恰一个 title；hero:* 0–max_heroes 个且限定词一致；panel:region ≤ 1 且不与 card 同页"""
    roles = defaultdict(list)
    errors = []

    for sh in slide.shapes:
        role, qual = parse_role(sh.name)
        if role not in ROLES:
            errors.append(f'未标记形状 {sh.name!r}（{sh.shape_type}）')
            continue
        if role == 'hero' and qual not in HERO_QUALIFIERS:
            errors.append(f'hero 限定词非法 {sh.name!r}，应为 hero:{"|".join(sorted(HERO_QUALIFIERS))}')
        if role == 'image' and qual not in IMAGE_QUALIFIERS:
            errors.append(f'image 限定词非法 {sh.name!r}，应为 image:{"|".join(sorted(IMAGE_QUALIFIERS))}')
        if role == 'panel' and qual not in PANEL_QUALIFIERS:
            errors.append(f'panel 限定词非法 {sh.name!r}，应为 panel 或 panel:region')
        if role == 'card' and not qual.isdigit():
            errors.append(f'card 需要序号 {sh.name!r}，应为 card:1、card:2 …')
        roles[role].append((sh, qual))

    n_title = len(roles['title'])
    n_hero = len(roles['hero'])
    if page_type in CONTENT_TYPES:
        if n_title != 1:
            errors.append(f'内容页应恰有一个 title，实际 {n_title}')
        if n_hero > max_heroes:
            errors.append(f'内容页 hero:* 最多 {max_heroes} 个，实际 {n_hero}')
        quals = {q for _, q in roles['hero']}
        if len(quals) > 1:
            errors.append(f'同页多个 hero:* 的限定词必须一致，实际 {sorted(quals)}')
        regions = [sh for sh, q in roles['panel'] if q == 'region']
        if len(regions) > 1:
            errors.append(f'panel:region（大容器）每页最多一个，实际 {len(regions)}')
        if regions and roles['card']:
            errors.append('panel:region 与 card:N 不同页出现（中 + 大容器不同页）')
    elif page_type in SPECIAL_TYPES:
        if n_title > 1:
            errors.append(f'特殊页最多一个 title，实际 {n_title}')
    else:
        errors.append(f'manifest 页类型未知：{page_type!r}')

    # 卡片序号应连续
    idx = sorted(int(q) for _, q in roles['card'])
    if idx and idx != list(range(1, len(idx) + 1)):
        errors.append(f'card 序号不连续：{idx}')

    return roles, errors


def load_manifest(path):
    """读 deck.manifest.yaml，返回 {page: type}"""
    import yaml
    with open(path, encoding='utf-8') as f:
        m = yaml.safe_load(f)
    out = {}
    for pg in m.get('pages', []):
        out[int(pg['page'])] = pg.get('type', 'content')
    return out


def main(path, manifest=None):
    from pptx import Presentation
    prs = Presentation(path)
    info = load_manifest(manifest) if manifest else {}
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        ptype = info.get(i, 'content')
        roles, errors = classify_slide(slide, ptype)
        inv = ', '.join(f'{r}×{len(v)}' for r, v in sorted(roles.items()) if v)
        print(f'[{i:02d}] {ptype:8s} {inv}')
        for e in errors:
            print(f'      ✗ {e}')
        total += len(errors)
    print(f'\n{"通过" if total == 0 else f"{total} 处违规"}')
    return total


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(1 if main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None) else 0)
