"""
roles.py — 形状角色标记校验（对应 skill 00「形状角色标记」）

用法：
    python roles.py deck.pptx [deck.manifest.yaml]   # 打印每页角色清单 + 违规；给 manifest 则按页类型与 hero=void 判定
    from roles import classify_slide       # 在 validate_design.py 里调用

classify_slide(slide, page_type) 返回 (roles, errors)：
    roles  : dict  role -> [ (shape, qualifier) ]   供其它检查按角色取形状
    errors : list  本页违规描述，空列表即通过
"""
import sys
from collections import defaultdict

ROLES = {
    'title', 'kicker', 'conclusion', 'body',
    'hero', 'card', 'panel', 'tag', 'divider', 'mask',
    'image', 'icon', 'chart', 'table', 'arrow',
    'source', 'pagenum', 'bg', 'deco',
}
HERO_QUALIFIERS = {'big-number', 'big-label', 'chart', 'table', 'image'}
IMAGE_QUALIFIERS = {'full-bleed', 'atmosphere', 'semantic', 'small'}
CONTENT_TYPES = {'content'}                 # 需要恰一个 title + 恰一个 hero 的页类型
SPECIAL_TYPES = {'cover', 'agenda', 'section', 'closing', 'quote'}


def parse_role(name: str):
    """'hero:big-number' -> ('hero', 'big-number'); 'title' -> ('title', '')"""
    role, _, qual = (name or '').partition(':')
    return role, qual


def classify_slide(slide, page_type: str = 'content', hero_void: bool = False):
    """hero_void=True 表示 manifest.hero == void：本页不得有 hero:* 形状"""
    roles = defaultdict(list)
    errors = []

    for sh in slide.shapes:
        role, qual = parse_role(sh.name)
        if role not in ROLES:
            errors.append(f'未标记形状 {sh.name!r}（{sh.shape_type}）')
            continue
        if role == 'hero' and qual not in HERO_QUALIFIERS:
            errors.append(f'hero 限定词非法 {sh.name!r}，应为 hero:{"|".join(sorted(HERO_QUALIFIERS))}')
        if role == 'image' and qual not in IMAGE_QUALIFIERS:
            errors.append(f'image 限定词非法 {sh.name!r}，应为 image:{"|".join(sorted(IMAGE_QUALIFIERS))}')
        if role == 'card' and not qual.isdigit():
            errors.append(f'card 需要序号 {sh.name!r}，应为 card:1、card:2 …')
        roles[role].append((sh, qual))

    n_title = len(roles['title'])
    n_hero = len(roles['hero'])
    if page_type in CONTENT_TYPES:
        if n_title != 1:
            errors.append(f'内容页应恰有一个 title，实际 {n_title}')
        want = 0 if hero_void else 1
        if n_hero != want:
            errors.append(f'内容页应有 {want} 个 hero:*（manifest.hero={"void" if hero_void else "non-void"}），实际 {n_hero}')
    elif page_type in SPECIAL_TYPES:
        if n_title > 1:
            errors.append(f'特殊页最多一个 title，实际 {n_title}')
    else:
        errors.append(f'manifest 页类型未知：{page_type!r}')

    # 卡片序号应连续
    idx = sorted(int(q) for _, q in roles['card'])
    if idx and idx != list(range(1, len(idx) + 1)):
        errors.append(f'card 序号不连续：{idx}')

    return roles, errors


def load_manifest(path):
    """读 deck.manifest.yaml，返回 {page: (type, hero_void)}"""
    import yaml
    with open(path, encoding='utf-8') as f:
        m = yaml.safe_load(f)
    out = {}
    for pg in m.get('pages', []):
        out[int(pg['page'])] = (pg.get('type', 'content'), pg.get('hero') == 'void')
    return out


def main(path, manifest=None):
    from pptx import Presentation
    prs = Presentation(path)
    info = load_manifest(manifest) if manifest else {}
    total = 0
    for i, slide in enumerate(prs.slides, 1):
        ptype, hero_void = info.get(i, ('content', False))
        roles, errors = classify_slide(slide, ptype, hero_void)
        inv = ', '.join(f'{r}×{len(v)}' for r, v in sorted(roles.items()) if v)
        print(f'[{i:02d}] {ptype:8s} {inv}')
        for e in errors:
            print(f'      ✗ {e}')
        total += len(errors)
    print(f'\n{"通过" if total == 0 else f"{total} 处违规"}')
    return total


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(1 if main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None) else 0)
